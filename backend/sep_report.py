from datetime import date
from datetime import datetime
from decimal import Decimal
from decimal import InvalidOperation
from math import isfinite
from pathlib import Path
import re

import xlsxwriter
from openpyxl import load_workbook


OUTPUT_HEADERS = (
    "تاریخ پرداخت",
    "نام مشتری",
    "شناسه راهکاران",
    "شماره راهکاران",
    "شماره موبایل",
    "مبلغ",
    "کد پرداخت",
    "وضعیت",
)

ADMIN_COLUMN_ALIASES = {
    "payment_date": (
        "تاریخ پرداخت",
        "تاريخ پرداخت",
    ),
    "customer_name": (
        "نام مشتری",
        "نام مشتري",
        "نام و نام خانوادگی مشتری",
        "نام و نام خانوادگي مشتري",
    ),
    "rahkaran_id": (
        "شناسه راهکاران",
        "شناسه راهکار",
    ),
    "rahkaran_number": (
        "شماره راهکاران",
        "شماره راهکار",
    ),
    "mobile_number": (
        "شماره موبایل",
        "شماره موبايل",
        "موبایل",
        "موبايل",
        "شماره همراه",
        "تلفن همراه",
    ),
    "amount": (
        "مبلغ",
        "مبلغ پرداخت",
        "مبلغ پرداختی",
    ),
    "payment_code": (
        "کد پرداخت",
        "شناسه پرداخت",
    ),
}

SEP_TRACKING_CODE_ALIASES = (
    "کد رهگیری",
    "کد رهگيری",
    "کد رهگيري",
    "شماره رهگیری",
    "شماره رهگيری",
    "شماره رهگيري",
    "کد پیگیری",
    "کد پيگيری",
    "کد پيگيري",
    "شماره پیگیری",
    "شماره پيگيری",
    "شماره پيگيري",
)

PERSIAN_DIGITS = str.maketrans(
    "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
    "01234567890123456789",
)

NUMERIC_CODE_PATTERN = re.compile(
    r"^[+-]?\d+(?:\.0+)?(?:[eE][+-]?\d+)?$"
)


def has_value(value):
    if value is None:
        return False

    if isinstance(value, str):
        return bool(value.strip())

    return True


def normalize_text(value):
    if value is None:
        return ""

    return (
        str(value)
        .translate(PERSIAN_DIGITS)
        .replace("ي", "ی")
        .replace("ك", "ک")
        .replace("ۀ", "ه")
        .replace("ة", "ه")
        .replace("\u200c", " ")
        .replace("\u200e", "")
        .replace("\u200f", "")
        .replace("\xa0", " ")
        .strip()
    )


def normalize_header(value):
    return re.sub(
        r"[\s_\-–—:/\\]+",
        "",
        normalize_text(value),
    ).casefold()


def normalize_code(value):
    if value is None or isinstance(value, bool):
        return ""

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float):
        if not isfinite(value):
            return ""

        if value.is_integer():
            return str(int(value))

        return format(value, ".15g")

    text = "".join(
        normalize_text(value).split()
    )

    if not text:
        return ""

    numeric_text = (
        text
        .replace(",", "")
        .replace("٬", "")
        .replace("،", "")
    )

    if NUMERIC_CODE_PATTERN.fullmatch(numeric_text):
        try:
            number = Decimal(numeric_text)

            if number == number.to_integral_value():
                return str(int(number))

        except InvalidOperation:
            pass

    return text.casefold()


def display_identifier(value):
    if value is None:
        return ""

    if isinstance(value, bool):
        return str(value)

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float):
        if not isfinite(value):
            return ""

        if value.is_integer():
            return str(int(value))

        return format(value, ".15g")

    return str(value).strip()


def find_column(headers, aliases):
    normalized_aliases = {
        normalize_header(alias)
        for alias in aliases
    }

    for index, header in enumerate(headers):
        if normalize_header(header) in normalized_aliases:
            return index

    return None


def find_admin_header(worksheet):
    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=1,
            max_row=50,
            values_only=True,
        ),
        start=1,
    ):
        indexes = {
            key: find_column(row, aliases)
            for key, aliases in ADMIN_COLUMN_ALIASES.items()
        }

        if all(
            index is not None
            for index in indexes.values()
        ):
            return row_number, indexes

    required_headers = "، ".join(
        aliases[0]
        for aliases in ADMIN_COLUMN_ALIASES.values()
    )

    raise ValueError(
        "ستون‌های موردنیاز در اکسل ادمین پیدا نشدند. "
        f"ستون‌های لازم: {required_headers}"
    )


def find_sep_header(worksheet):
    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=1,
            max_row=50,
            values_only=True,
        ),
        start=1,
    ):
        tracking_index = find_column(
            row,
            SEP_TRACKING_CODE_ALIASES,
        )

        if tracking_index is not None:
            return row_number, tracking_index

    raise ValueError(
        "ستون «کد رهگیری» در اکسل SEP پیدا نشد."
    )


def value_at(row, index):
    if index < len(row):
        return row[index]

    return None


def read_admin_rows(file_path):
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook.active

    try:
        header_row, indexes = find_admin_header(
            worksheet
        )

        rows = []

        for row in worksheet.iter_rows(
            min_row=header_row + 1,
            values_only=True,
        ):
            selected_row = (
                value_at(row, indexes["payment_date"]),
                value_at(row, indexes["customer_name"]),
                value_at(row, indexes["rahkaran_id"]),
                value_at(row, indexes["rahkaran_number"]),
                value_at(row, indexes["mobile_number"]),
                value_at(row, indexes["amount"]),
                value_at(row, indexes["payment_code"]),
            )

            if any(
                has_value(value)
                for value in selected_row
            ):
                rows.append(selected_row)

        return rows

    finally:
        workbook.close()


def read_sep_tracking_codes(file_path):
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook.active

    try:
        header_row, tracking_index = find_sep_header(
            worksheet
        )

        tracking_codes = set()

        for row in worksheet.iter_rows(
            min_row=header_row + 1,
            values_only=True,
        ):
            tracking_code = normalize_code(
                value_at(row, tracking_index)
            )

            if tracking_code:
                tracking_codes.add(tracking_code)

        return tracking_codes

    finally:
        workbook.close()


def write_regular_value(
    worksheet,
    row_number,
    column_number,
    value,
    cell_format,
):
    if value is None:
        worksheet.write_blank(
            row_number,
            column_number,
            None,
            cell_format,
        )
        return

    worksheet.write(
        row_number,
        column_number,
        value,
        cell_format,
    )


def write_date_value(
    worksheet,
    row_number,
    column_number,
    value,
    date_format,
    text_format,
):
    if isinstance(value, (datetime, date)):
        worksheet.write_datetime(
            row_number,
            column_number,
            value,
            date_format,
        )
        return

    write_regular_value(
        worksheet,
        row_number,
        column_number,
        value,
        text_format,
    )


def create_sep_comparison_report(
    admin_file_path,
    sep_file_path,
    output_path,
):
    admin_file_path = Path(admin_file_path)
    sep_file_path = Path(sep_file_path)
    output_path = Path(output_path)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    admin_rows = read_admin_rows(
        admin_file_path
    )

    sep_tracking_codes = read_sep_tracking_codes(
        sep_file_path
    )

    workbook = xlsxwriter.Workbook(
        str(output_path),
        {
            "constant_memory": True,
            "strings_to_numbers": False,
            "strings_to_urls": False,
        },
    )

    try:
        worksheet = workbook.add_worksheet(
            "تطبیق SEP"
        )

        worksheet.right_to_left()
        worksheet.hide_gridlines(2)
        worksheet.freeze_panes(1, 0)
        worksheet.set_default_row(21)

        header_format = workbook.add_format({
            "font_name": "Peyda",
            "font_size": 11,
            "bold": True,
            "font_color": "#FFFFFF",
            "bg_color": "#2F75B5",
            "align": "center",
            "valign": "vcenter",
            "border": 1,
            "border_color": "#D9E2F3",
        })

        base_format = {
            "font_name": "Peyda",
            "font_size": 11,
            "valign": "vcenter",
            "bottom": 1,
            "bottom_color": "#E7EAF0",
        }

        text_format = workbook.add_format({
            **base_format,
            "align": "right",
        })

        center_format = workbook.add_format({
            **base_format,
            "align": "center",
        })

        date_format = workbook.add_format({
            **base_format,
            "align": "center",
            "num_format": "yyyy-mm-dd",
        })

        amount_format = workbook.add_format({
            **base_format,
            "align": "center",
            "num_format": "#,##0",
        })

        code_format = workbook.add_format({
            **base_format,
            "align": "center",
            "num_format": "@",
        })

        ok_format = workbook.add_format({
            **base_format,
            "bold": True,
            "font_color": "#375623",
            "bg_color": "#E2F0D9",
            "align": "center",
        })

        no_format = workbook.add_format({
            **base_format,
            "bold": True,
            "font_color": "#9C0006",
            "bg_color": "#F4CCCC",
            "align": "center",
        })

        for column_number, header in enumerate(
            OUTPUT_HEADERS
        ):
            worksheet.write(
                0,
                column_number,
                header,
                header_format,
            )

        for row_index, admin_row in enumerate(
            admin_rows,
            start=1,
        ):
            write_date_value(
                worksheet,
                row_index,
                0,
                admin_row[0],
                date_format,
                center_format,
            )

            write_regular_value(
                worksheet,
                row_index,
                1,
                admin_row[1],
                text_format,
            )

            for column_number in (2, 3, 4):
                worksheet.write_string(
                    row_index,
                    column_number,
                    display_identifier(
                        admin_row[column_number]
                    ),
                    code_format,
                )

            write_regular_value(
                worksheet,
                row_index,
                5,
                admin_row[5],
                amount_format,
            )

            payment_code = admin_row[6]
            payment_key = normalize_code(
                payment_code
            )

            worksheet.write_string(
                row_index,
                6,
                display_identifier(payment_code),
                code_format,
            )

            is_matched = bool(
                payment_key
                and payment_key in sep_tracking_codes
            )

            worksheet.write_string(
                row_index,
                7,
                "OK" if is_matched else "NO",
                ok_format if is_matched else no_format,
            )

        worksheet.set_row(0, 26)
        worksheet.set_column(0, 0, 18)
        worksheet.set_column(1, 1, 28)
        worksheet.set_column(2, 3, 20)
        worksheet.set_column(4, 4, 18)
        worksheet.set_column(5, 5, 20)
        worksheet.set_column(6, 6, 24)
        worksheet.set_column(7, 7, 12)

        worksheet.autofilter(
            0,
            0,
            max(len(admin_rows), 1),
            len(OUTPUT_HEADERS) - 1,
        )

    finally:
        workbook.close()

    return str(output_path)
