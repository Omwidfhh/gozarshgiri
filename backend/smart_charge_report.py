import re
import unicodedata
from copy import copy
from datetime import date, datetime
from math import isfinite
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

TITLE_HEADER_ALIASES = (
    "عنوان کالا",
    "نام کالا",
    "عنوان محصول",
    "نام محصول",
)

CODE_HEADER_ALIASES = (
    "کد کالا",
    "کد محصول",
)

TECHNICAL_CODE_HEADER_ALIASES = (
    "مشخصه فنی کالا",
    "مشخصات فنی کالا",
    "مشخصه فنی",
    "مشخصات فنی",
    "بارکد",
)

SITE_STOCK_HEADER_ALIASES = (
    "موجودی قابل فروش",
    "موجودي قابل فروش",
)

PRODUCT_STOCK_HEADER_ALIASES = (
    "موجودی انبار محصول",
    "موجودي انبار محصول",
)

BARCODE_HEADER_ALIASES = (
    "بارکد",
    "بار کد",
    "کد بارکد",
)

DEFINITION_DATE_HEADER_ALIASES = (
    "تاریخ تعریف",
    "تاريخ تعريف",
    "تاریخ ثبت",
    "تاريخ ثبت",
)

STATUS_HEADER = "وضعیت بررسی"
DEFINITION_DATE_HEADER = "تاریخ تعریف"

CHARGE_STATUS = "پیشنهاد شارژ سایت"
DUPLICATE_STATUS = "تکراری موجود در سایت"
SIMILAR_ACTIVE_STATUS = "کد مشابه در سایت موجود است"
ALTERNATIVE_SELECTED_STATUS = "کد مشابه دیگری برای شارژ انتخاب شد"
UNKNOWN_DUPLICATE_DATE_STATUS = "تکراری؛ تاریخ تعریف نامشخص"

SOFT_RED_FILL = PatternFill(
    fill_type="solid",
    fgColor="F4CCCC",
)

SOFT_BLUE_FILL = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7",
)

PEYDA_FONT = Font(
    name="Peyda",
    size=11,
)

PEYDA_HEADER_FONT = Font(
    name="Peyda",
    size=11,
    bold=True,
)

DIGIT_TRANSLATION = str.maketrans(
    "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
    "01234567890123456789",
)

# کد محصول داخل عنوان معمولاً حداقل چهار رقم دارد. اعداد قد،
# سایز و تعداد پک کوتاه‌ترند و عمداً در کلید تکراری باقی می‌مانند.
EMBEDDED_PRODUCT_CODE_PATTERN = re.compile(
    r"(?<!\d)\d{4,}(?!\d)"
)

SPACE_PATTERN = re.compile(r"\s+")

# بعضی خروجی‌ها یک X جدا با فاصله قبل و بعد دارند که فقط
# علامت نگارشی/جداکننده است. X داخل XL، 2X، X7 یا X-Large
# عمداً حذف نمی‌شود.
STANDALONE_X_PATTERN = re.compile(
    r"(?<![\w/+\-])x(?![\w/+\-])",
    re.IGNORECASE,
)

# وجود یا نبود پرانتز دور Normal، ECO و اندازه‌ها نباید
# باعث ساخته‌شدن یک گروه محصول جدا شود.
ROUND_PARENTHESES_PATTERN = re.compile(
    r"[()（）]"
)

DATE_PART_PATTERN = re.compile(r"\d+")

BUTTON_ONE_ORANGE = "FCE4D6"


def normalize_text(value):
    if value is None:
        return ""

    text = unicodedata.normalize(
        "NFKC",
        str(value),
    )
    text = "".join(
        character
        for character in text
        if unicodedata.category(character) != "Cf"
    )

    return SPACE_PATTERN.sub(
        " ",
        text
        .translate(DIGIT_TRANSLATION)
        .replace("ي", "ی")
        .replace("ك", "ک")
        .replace("\u200c", " ")
        .replace("\u200e", "")
        .replace("\u200f", "")
        .replace("\xa0", " ")
        .strip(),
    )


def normalize_header(value):
    return "".join(
        normalize_text(value).split()
    ).casefold()


def normalize_code(value):
    if value is None:
        return ""

    if isinstance(value, bool):
        return str(value)

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float):
        if not isfinite(value):
            return ""

        if value.is_integer():
            return str(int(value))

        return format(value, ".15g")

    return "".join(
        normalize_text(value).split()
    )


def to_number(value):
    if value is None or isinstance(value, bool):
        return None

    if isinstance(value, (int, float)):
        number = float(value)

        if isfinite(number):
            return number

        return None

    text = (
        normalize_text(value)
        .replace(",", "")
        .replace("٬", "")
        .replace(" ", "")
    )

    if not text:
        return None

    try:
        number = float(text)
    except ValueError:
        return None

    if not isfinite(number):
        return None

    return number


def definition_date_sort_key(
    value,
    excel_epoch=None,
):
    if isinstance(value, datetime):
        return (
            value.year,
            value.month,
            value.day,
            value.hour,
            value.minute,
            value.second,
            value.microsecond,
        )

    if isinstance(value, date):
        return (
            value.year,
            value.month,
            value.day,
            0,
            0,
            0,
            0,
        )

    if (
        isinstance(value, (int, float))
        and not isinstance(value, bool)
        and isfinite(float(value))
        and 1 <= float(value) <= 100000
    ):
        try:
            converted = from_excel(
                value,
                epoch=excel_epoch,
            )
        except (TypeError, ValueError, OverflowError):
            converted = None

        if isinstance(converted, datetime):
            return definition_date_sort_key(
                converted
            )

        if isinstance(converted, date):
            return definition_date_sort_key(
                converted
            )

    text = normalize_text(value)

    if not text or text.casefold() in {
        "n/a",
        "na",
        "none",
        "null",
        "-",
    }:
        return None

    date_parts = [
        int(part)
        for part in DATE_PART_PATTERN.findall(text)
    ]

    if len(date_parts) == 1:
        compact_date = str(date_parts[0])

        if len(compact_date) == 8:
            date_parts = [
                int(compact_date[0:4]),
                int(compact_date[4:6]),
                int(compact_date[6:8]),
            ]

    if len(date_parts) < 3:
        return None

    first, second, third = date_parts[:3]

    if first >= 1000:
        year, month, day = first, second, third
    elif third >= 1000:
        year, month, day = third, second, first
    else:
        return None

    if not (
        1 <= month <= 12
        and 1 <= day <= 31
    ):
        return None

    time_parts = date_parts[3:6]
    hour = time_parts[0] if len(time_parts) >= 1 else 0
    minute = time_parts[1] if len(time_parts) >= 2 else 0
    second_value = time_parts[2] if len(time_parts) >= 3 else 0

    if not (
        0 <= hour <= 23
        and 0 <= minute <= 59
        and 0 <= second_value <= 59
    ):
        hour = 0
        minute = 0
        second_value = 0

    return (
        year,
        month,
        day,
        hour,
        minute,
        second_value,
        0,
    )


def duplicate_title_key(
    value,
    technical_code=None,
):
    title = normalize_text(value).casefold()

    if not title:
        return ""

    normalized_technical_code = normalize_code(
        technical_code
    ).casefold()
    title_without_code = title

    if normalized_technical_code:
        exact_code_pattern = re.compile(
            rf"(?<!\d)"
            rf"{re.escape(normalized_technical_code)}"
            rf"(?!\d)"
        )
        title_without_code, _ = (
            exact_code_pattern.subn(
                " ",
                title_without_code,
            )
        )

    # الگوی عمومی فقط برای ردیفی استفاده می‌شود که اصلاً کد فنی
    # ندارد. اگر کد فنی موجود ولی با عنوان ناسازگار باشد، عدد دیگری
    # را حذف نمی‌کنیم؛ آن ردیف جدا می‌ماند تا با محصول نامرتبط یکی نشود.
    if not normalized_technical_code:
        title_without_code = (
            EMBEDDED_PRODUCT_CODE_PATTERN.sub(
                " ",
                title_without_code,
            )
        )

    title_without_format_markers = (
        STANDALONE_X_PATTERN.sub(
            " ",
            title_without_code,
        )
    )
    title_without_format_markers = (
        ROUND_PARENTHESES_PATTERN.sub(
            "",
            title_without_format_markers,
        )
    )

    # فاصله، تب و فاصله اطراف | یا کروشه نباید یک محصول را
    # به گروه‌های جدا تبدیل کند. X جدا و پرانتز نیز فقط تفاوت
    # نگارشی محسوب می‌شوند، نه تفاوت واقعی محصول.
    return SPACE_PATTERN.sub(
        "",
        title_without_format_markers,
    ).replace(
        "ـ",
        "",
    ).strip(" -–—_|،,:؛")


def find_column(worksheet, aliases):
    wanted = {
        normalize_header(alias)
        for alias in aliases
    }

    try:
        header_row = next(
            worksheet.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True,
            )
        )
    except StopIteration:
        return None

    for column_number, header in enumerate(
        header_row,
        start=1,
    ):
        if normalize_header(header) in wanted:
            return column_number

    return None


def require_column(worksheet, aliases, readable_name):
    column_number = find_column(
        worksheet,
        aliases,
    )

    if column_number is None:
        raise ValueError(
            f"ستون «{readable_name}» در خروجی پیدا نشد."
        )

    return column_number


def read_definition_dates(file_path):
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )
    worksheet = workbook.active

    try:
        barcode_column = find_column(
            worksheet,
            BARCODE_HEADER_ALIASES,
        )
        definition_date_column = find_column(
            worksheet,
            DEFINITION_DATE_HEADER_ALIASES,
        )

        if barcode_column is None:
            raise ValueError(
                "ستون «بارکد» در اکسل سوم پیدا نشد."
            )

        if definition_date_column is None:
            raise ValueError(
                "ستون «تاریخ تعریف» در اکسل سوم پیدا نشد."
            )

        required_column = max(
            barcode_column,
            definition_date_column,
        )
        definition_dates = {}

        for row in worksheet.iter_rows(
            min_row=2,
            values_only=False,
        ):
            if len(row) < required_column:
                continue

            barcode = normalize_code(
                row[barcode_column - 1].value
            )

            if not barcode:
                continue

            date_cell = row[
                definition_date_column - 1
            ]
            candidate = {
                "value": date_cell.value,
                "number_format": date_cell.number_format,
                "sort_key": definition_date_sort_key(
                    date_cell.value,
                    workbook.epoch,
                ),
            }
            existing = definition_dates.get(
                barcode
            )

            if existing is None:
                definition_dates[barcode] = candidate
                continue

            candidate_key = candidate["sort_key"]
            existing_key = existing["sort_key"]

            if (
                candidate_key is not None
                and (
                    existing_key is None
                    or candidate_key < existing_key
                )
            ):
                definition_dates[barcode] = candidate

        return definition_dates

    finally:
        workbook.close()


def copy_cell_style(source_cell, target_cell):
    if source_cell.has_style:
        target_cell._style = copy(
            source_cell._style
        )

    target_cell.alignment = copy(
        source_cell.alignment
    )
    target_cell.border = copy(
        source_cell.border
    )
    target_cell.fill = copy(
        source_cell.fill
    )
    target_cell.protection = copy(
        source_cell.protection
    )
    target_cell.number_format = (
        source_cell.number_format
    )


def prepare_status_column(worksheet):
    existing_column = find_column(
        worksheet,
        (STATUS_HEADER,),
    )

    if existing_column is not None:
        return existing_column

    last_header_column = max(
        (
            cell.column
            for cell in worksheet[1]
            if cell.value not in (None, "")
        ),
        default=worksheet.max_column,
    )
    status_column = last_header_column + 1

    style_source = worksheet.cell(
        row=1,
        column=max(last_header_column, 1),
    )
    status_cell = worksheet.cell(
        row=1,
        column=status_column,
    )

    if not isinstance(style_source, MergedCell):
        copy_cell_style(
            style_source,
            status_cell,
        )

    status_cell.value = STATUS_HEADER
    status_cell.font = copy(PEYDA_HEADER_FONT)

    return status_column


def prepare_definition_date_column(worksheet):
    existing_column = find_column(
        worksheet,
        DEFINITION_DATE_HEADER_ALIASES,
    )

    if existing_column is not None:
        return existing_column

    last_header_column = max(
        (
            cell.column
            for cell in worksheet[1]
            if cell.value not in (None, "")
        ),
        default=worksheet.max_column,
    )
    date_column = last_header_column + 1
    style_source = worksheet.cell(
        row=1,
        column=max(last_header_column, 1),
    )
    date_header_cell = worksheet.cell(
        row=1,
        column=date_column,
    )

    if not isinstance(style_source, MergedCell):
        copy_cell_style(
            style_source,
            date_header_cell,
        )

    date_header_cell.value = DEFINITION_DATE_HEADER
    date_header_cell.font = copy(PEYDA_HEADER_FONT)

    return date_column


def write_definition_date_cell(
    worksheet,
    row_number,
    date_column,
    matched_definition,
):
    output_cell = worksheet.cell(
        row=row_number,
        column=date_column,
    )
    style_source = worksheet.cell(
        row=row_number,
        column=max(date_column - 1, 1),
    )

    if not isinstance(style_source, MergedCell):
        copy_cell_style(
            style_source,
            output_cell,
        )

    output_cell.value = None
    output_cell.font = copy(PEYDA_FONT)

    if matched_definition is None:
        return

    output_cell.value = matched_definition[
        "value"
    ]

    if matched_definition["number_format"]:
        output_cell.number_format = (
            matched_definition["number_format"]
        )


def fill_entire_row(
    worksheet,
    row_number,
    max_column,
    fill,
):
    for column_number in range(
        1,
        max_column + 1,
    ):
        cell = worksheet.cell(
            row=row_number,
            column=column_number,
        )

        if isinstance(cell, MergedCell):
            continue

        cell.fill = copy(fill)


def clear_previous_analysis_fills(worksheet):
    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        for column_number in range(
            1,
            worksheet.max_column + 1,
        ):
            cell = worksheet.cell(
                row=row_number,
                column=column_number,
            )

            if isinstance(cell, MergedCell):
                continue

            fill = cell.fill
            color = fill.fgColor
            color_value = (
                color.rgb
                if color.type == "rgb"
                else None
            )

            if (
                fill.fill_type == "solid"
                and color_value
                and any(
                    str(color_value).upper().endswith(
                        color_code
                    )
                    for color_code in (
                        BUTTON_ONE_ORANGE,
                        "F4CCCC",
                        "D9EAF7",
                    )
                )
            ):
                cell.fill = PatternFill()


def apply_peyda_font(worksheet):
    for row in worksheet.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue

            old_font = cell.font
            cell.font = Font(
                name="Peyda",
                size=11,
                bold=old_font.bold,
                italic=old_font.italic,
                color=copy(old_font.color),
                underline=old_font.underline,
                strike=old_font.strike,
            )


def analyze_smart_charge_output(
    output_path,
    definition_file,
):
    output_path = Path(output_path)
    definition_dates = read_definition_dates(
        definition_file
    )

    workbook = load_workbook(
        filename=output_path,
        read_only=False,
        data_only=False,
    )
    worksheet = workbook.active

    try:
        title_column = require_column(
            worksheet,
            TITLE_HEADER_ALIASES,
            "عنوان کالا",
        )
        technical_code_column = require_column(
            worksheet,
            TECHNICAL_CODE_HEADER_ALIASES,
            "مشخصه فنی کالا",
        )
        site_stock_column = require_column(
            worksheet,
            SITE_STOCK_HEADER_ALIASES,
            "موجودی قابل فروش",
        )
        product_stock_column = require_column(
            worksheet,
            PRODUCT_STOCK_HEADER_ALIASES,
            "موجودی انبار محصول",
        )
        status_column = prepare_status_column(
            worksheet
        )
        definition_date_column = (
            prepare_definition_date_column(
                worksheet
            )
        )
        last_output_column = max(
            status_column,
            definition_date_column,
        )

        title_groups = {}
        matched_definition_count = 0

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):
            site_stock = to_number(
                worksheet.cell(
                    row=row_number,
                    column=site_stock_column,
                ).value
            )

            product_code = normalize_code(
                worksheet.cell(
                    row=row_number,
                    column=technical_code_column,
                ).value
            )
            matched_definition = (
                definition_dates.get(
                    product_code
                )
            )

            if matched_definition is not None:
                matched_definition_count += 1

            write_definition_date_cell(
                worksheet=worksheet,
                row_number=row_number,
                date_column=definition_date_column,
                matched_definition=matched_definition,
            )

            title_key = duplicate_title_key(
                worksheet.cell(
                    row=row_number,
                    column=title_column,
                ).value,
                product_code,
            )

            if not title_key:
                title_key = f"__row_{row_number}"

            product_stock = to_number(
                worksheet.cell(
                    row=row_number,
                    column=product_stock_column,
                ).value
            )

            group = title_groups.setdefault(
                title_key,
                [],
            )
            group.append({
                "row": row_number,
                "code": product_code,
                "site_stock": site_stock,
                "product_stock": product_stock,
                "definition_sort_key": (
                    matched_definition["sort_key"]
                    if matched_definition is not None
                    else None
                ),
            })

        duplicate_rows = set()
        charge_rows = set()
        row_statuses = {}

        for group in title_groups.values():
            active_rows = [
                item
                for item in group
                if (
                    item["site_stock"] is not None
                    and item["site_stock"] > 0
                )
            ]
            available_candidates = [
                item
                for item in group
                if (
                    item["site_stock"] == 0
                    and item["product_stock"] is not None
                    and item["product_stock"] > 0
                )
            ]

            if len(active_rows) >= 2:
                # از بین کدهای فعال تکراری فقط قدیمی‌ترین تاریخ تعریف
                # آبی می‌شود. تاریخ خالی یا نامعتبر در انتخاب قدیمی‌ترین
                # شرکت نمی‌کند و در تساوی، ردیف بالاتر انتخاب می‌شود.
                dated_active_rows = [
                    item
                    for item in active_rows
                    if item["definition_sort_key"]
                    is not None
                ]

                if dated_active_rows:
                    oldest_active_row = min(
                        dated_active_rows,
                        key=lambda item: (
                            item[
                                "definition_sort_key"
                            ],
                            item["row"],
                        ),
                    )
                    duplicate_rows.add(
                        oldest_active_row["row"]
                    )
                else:
                    for item in active_rows:
                        row_statuses[item["row"]] = (
                            UNKNOWN_DUPLICATE_DATE_STATUS
                        )

            if active_rows:
                for item in available_candidates:
                    row_statuses[item["row"]] = (
                        SIMILAR_ACTIVE_STATUS
                    )

                continue

            if not available_candidates:
                continue

            # وقتی هیچ کد مشابهی در سایت موجود نیست، فقط یک کد
            # پیشنهاد می‌شود؛ کدی که بیشترین موجودی انبار محصول دارد.
            selected_candidate = max(
                available_candidates,
                key=lambda item: (
                    item["product_stock"],
                    -item["row"],
                ),
            )
            charge_rows.add(
                selected_candidate["row"]
            )

            for item in available_candidates:
                if item["row"] == selected_candidate["row"]:
                    continue

                row_statuses[item["row"]] = (
                    ALTERNATIVE_SELECTED_STATUS
                )

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):
            status_cell = worksheet.cell(
                row=row_number,
                column=status_column,
            )
            status_cell.font = copy(PEYDA_FONT)

            if row_number in charge_rows:
                status_cell.value = CHARGE_STATUS
                fill_entire_row(
                    worksheet,
                    row_number,
                    last_output_column,
                    SOFT_RED_FILL,
                )
                continue

            if row_number in duplicate_rows:
                status_cell.value = DUPLICATE_STATUS
                fill_entire_row(
                    worksheet,
                    row_number,
                    last_output_column,
                    SOFT_BLUE_FILL,
                )
                continue

            status_cell.value = row_statuses.get(
                row_number
            )

        worksheet.column_dimensions[
            get_column_letter(status_column)
        ].width = 24
        worksheet.column_dimensions[
            get_column_letter(
                definition_date_column
            )
        ].width = 18
        worksheet.column_dimensions[
            get_column_letter(
                definition_date_column
            )
        ].bestFit = True
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = (
            f"A1:"
            f"{get_column_letter(worksheet.max_column)}"
            f"{max(worksheet.max_row, 1)}"
        )
        worksheet.sheet_view.rightToLeft = True

        workbook.save(output_path)

        return {
            "output_path": str(output_path),
            "charge_candidate_count": len(charge_rows),
            "duplicate_row_count": len(duplicate_rows),
            "definition_date_matched_count": (
                matched_definition_count
            ),
        }

    finally:
        workbook.close()


def create_smart_charge_report(
    site_file,
    product_stock_file,
    definition_file,
    output_path,
):
    # ایمپورت در زمان اجرا انجام می‌شود تا تحلیل خروجی به‌صورت
    # مستقل نیز قابل آزمایش و استفاده باشد.
    from report import create_report

    output_path = Path(output_path)
    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    # تمام پاک‌سازی‌ها و تطبیق موجودی دکمه اول ابتدا اجرا می‌شود.
    create_report(
        site_file,
        product_stock_file,
        output_path,
        highlight_zero_rows=False,
    )

    return analyze_smart_charge_output(
        output_path,
        definition_file,
    )
