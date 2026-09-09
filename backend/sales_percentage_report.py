from __future__ import annotations

from collections import Counter
from copy import copy
from datetime import date, datetime, timedelta
import math
from pathlib import Path
import re
import traceback
from typing import Iterable, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    from python_calamine import CalamineWorkbook
except Exception:  # pragma: no cover - only relevant if dependency is missing
    CalamineWorkbook = None


BLOCKED_ISSUERS = (
    "ابوالفضل فلاح",
    "رقیه پیرانی",
    "مهدیه تن ساز",
)

RED_FILL = PatternFill(fill_type="solid", fgColor="FFFFC7CE")
MAX_HEADER_SCAN_ROWS = 60
MAX_HEADER_SCAN_COLS = 250


def _fa_digits_to_en(value: str) -> str:
    return value.translate(
        str.maketrans(
            "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
            "01234567890123456789",
        )
    )


def _norm_text(value) -> str:
    if value is None:
        return ""
    text = _fa_digits_to_en(str(value))
    text = (
        text.replace("ي", "ی")
        .replace("ك", "ک")
        .replace("ة", "ه")
        .replace("ۀ", "ه")
        .replace("\u200c", " ")
        .replace("\u200e", " ")
        .replace("\u200f", " ")
        .replace("\ufeff", " ")
        .replace("\u00a0", " ")
    )
    return re.sub(r"\s+", " ", text).strip().lower()


def _compact_text(value) -> str:
    return re.sub(r"[\s_\-]+", "", _norm_text(value))


def _norm_key(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if value.is_integer():
            return str(int(value))
        return format(value, ".15g")

    text = _fa_digits_to_en(str(value)).strip()
    text = (
        text.replace("\u200c", "")
        .replace("\u200e", "")
        .replace("\u200f", "")
        .replace("\u00a0", "")
    )
    text = re.sub(r"\s+", "", text)
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".", 1)[0]
    return text


def _to_number(value) -> Optional[float]:
    if value is None or value == "":
        return 0.0
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return float(value)

    text = _fa_digits_to_en(str(value)).strip()
    text = (
        text.replace(",", "")
        .replace("٬", "")
        .replace("٫", ".")
        .replace("%", "")
        .replace("٪", "")
        .strip()
    )
    if not text:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return None


def _gregorian_to_jalali(gy: int, gm: int, gd: int) -> tuple[int, int, int]:
    g_days_in_month = [0, 31, 59, 90, 120, 151, 181, 212, 243, 273, 304, 334]
    gy2 = gy + 1 if gm > 2 else gy
    days = (
        355666
        + (365 * gy)
        + ((gy2 + 3) // 4)
        - ((gy2 + 99) // 100)
        + ((gy2 + 399) // 400)
        + gd
        + g_days_in_month[gm - 1]
    )
    jy = -1595 + 33 * (days // 12053)
    days %= 12053
    jy += 4 * (days // 1461)
    days %= 1461
    if days > 365:
        jy += (days - 1) // 365
        days = (days - 1) % 365
    if days < 186:
        jm = 1 + days // 31
        jd = 1 + days % 31
    else:
        jm = 7 + (days - 186) // 30
        jd = 1 + (days - 186) % 30
    return jy, jm, jd


def _jalali_to_gregorian(jy: int, jm: int, jd: int) -> tuple[int, int, int]:
    jy += 1595
    days = (
        -355668
        + (365 * jy)
        + ((jy // 33) * 8)
        + (((jy % 33) + 3) // 4)
        + jd
    )
    if jm < 7:
        days += (jm - 1) * 31
    else:
        days += ((jm - 7) * 30) + 186

    gy = 400 * (days // 146097)
    days %= 146097
    if days > 36524:
        days -= 1
        gy += 100 * (days // 36524)
        days %= 36524
        if days >= 365:
            days += 1
    gy += 4 * (days // 1461)
    days %= 1461
    if days > 365:
        gy += (days - 1) // 365
        days = (days - 1) % 365

    gd = days + 1
    month_lengths = [0, 31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    if (gy % 4 == 0 and gy % 100 != 0) or gy % 400 == 0:
        month_lengths[2] = 29
    gm = 1
    while gm <= 12 and gd > month_lengths[gm]:
        gd -= month_lengths[gm]
        gm += 1
    return gy, gm, gd


def _parse_jalali(value) -> Optional[tuple[int, int, int]]:
    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return _gregorian_to_jalali(value.year, value.month, value.day)

    if isinstance(value, float) and value.is_integer():
        value = int(value)

    text = _fa_digits_to_en(str(value)).strip()

    # Jalali: 1405/06/16, 1405-6-16, 1405.06.16, 14050616
    match = re.search(
        r"(?<!\d)(1[34]\d{2})\s*[/\-.]\s*(\d{1,2})\s*[/\-.]\s*(\d{1,2})(?!\d)",
        text,
    )
    if not match:
        match = re.search(r"(?<!\d)(1[34]\d{2})(\d{2})(\d{2})(?!\d)", text)

    if match:
        y, mo, d = map(int, match.groups())
        if not (1 <= mo <= 12 and 1 <= d <= 31):
            return None
        try:
            gy, gm, gd = _jalali_to_gregorian(y, mo, d)
            if _gregorian_to_jalali(gy, gm, gd) == (y, mo, d):
                return y, mo, d
        except Exception:
            return None

    # If Excel exported a Gregorian date as text, convert it to Jalali too.
    gregorian = re.search(
        r"(?<!\d)(20\d{2})\s*[/\-.]\s*(\d{1,2})\s*[/\-.]\s*(\d{1,2})(?!\d)",
        text,
    )
    if gregorian:
        gy, gm, gd = map(int, gregorian.groups())
        try:
            date(gy, gm, gd)
            return _gregorian_to_jalali(gy, gm, gd)
        except Exception:
            return None

    return None


def _jalali_next_day(value: tuple[int, int, int]) -> tuple[int, int, int]:
    gy, gm, gd = _jalali_to_gregorian(*value)
    next_g = date(gy, gm, gd) + timedelta(days=1)
    return _gregorian_to_jalali(next_g.year, next_g.month, next_g.day)


def _aliases_match(value, aliases: Iterable[str]) -> bool:
    value_norm = _norm_text(value)
    value_compact = _compact_text(value)
    if not value_norm:
        return False
    for alias in aliases:
        alias_norm = _norm_text(alias)
        alias_compact = _compact_text(alias)
        if value_norm == alias_norm or value_compact == alias_compact:
            return True
        if alias_norm in value_norm or value_norm in alias_norm:
            return True
    return False


def _find_header_row_ws(ws, groups: list[Iterable[str]]) -> int:
    best_row = None
    best_score = -1
    max_row = min(ws.max_row, MAX_HEADER_SCAN_ROWS)
    max_col = min(ws.max_column, MAX_HEADER_SCAN_COLS)
    for row in range(1, max_row + 1):
        values = [ws.cell(row, col).value for col in range(1, max_col + 1)]
        score = sum(any(_aliases_match(v, aliases) for v in values) for aliases in groups)
        if score > best_score:
            best_score = score
            best_row = row
    if best_row is None or best_score <= 0:
        raise ValueError("ردیف عنوان ستون‌ها پیدا نشد.")
    return best_row


def _find_col_ws(ws, header_row: int, aliases: Iterable[str], required: bool = True) -> Optional[int]:
    max_col = min(ws.max_column, MAX_HEADER_SCAN_COLS)
    for col in range(1, max_col + 1):
        if _aliases_match(ws.cell(header_row, col).value, aliases):
            return col
    if required:
        visible = [
            str(ws.cell(header_row, col).value)
            for col in range(1, max_col + 1)
            if ws.cell(header_row, col).value not in (None, "")
        ]
        raise ValueError(
            f"ستون «{next(iter(aliases))}» پیدا نشد. ستون‌های موجود: {', '.join(visible[:60])}"
        )
    return None


def _find_header_row_matrix(rows: list[list], groups: list[Iterable[str]]) -> int:
    best_index = None
    best_score = -1
    for index, row in enumerate(rows[:MAX_HEADER_SCAN_ROWS]):
        values = list(row[:MAX_HEADER_SCAN_COLS])
        score = sum(any(_aliases_match(v, aliases) for v in values) for aliases in groups)
        if score > best_score:
            best_score = score
            best_index = index
    if best_index is None or best_score <= 0:
        raise ValueError("ردیف عنوان ستون‌ها در فایل انتقال پیدا نشد.")
    return best_index


def _find_col_matrix(row: list, aliases: Iterable[str], required: bool = True) -> Optional[int]:
    for index, value in enumerate(row[:MAX_HEADER_SCAN_COLS]):
        if _aliases_match(value, aliases):
            return index
    if required:
        visible = [str(v) for v in row[:MAX_HEADER_SCAN_COLS] if v not in (None, "")]
        raise ValueError(
            f"ستون «{next(iter(aliases))}» در فایل انتقال پیدا نشد. ستون‌های موجود: {', '.join(visible[:60])}"
        )
    return None


def _get_row_value(row: list, index: int):
    return row[index] if index < len(row) else None


def _issuer_is_blocked(value) -> bool:
    text = _compact_text(value)
    if not text:
        return False
    return any(_compact_text(name) in text for name in BLOCKED_ISSUERS)


def _copy_cell_style(src, dst) -> None:
    try:
        if src.has_style:
            dst._style = copy(src._style)
        dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
    except Exception:
        pass


def _load_calamine_rows(path: Path) -> list[list]:
    if CalamineWorkbook is None:
        raise RuntimeError(
            "کتابخانه python-calamine نصب نیست. داخل Codespaces بزن: pip install -r requirements.txt"
        )
    book = CalamineWorkbook.from_path(str(path))
    sheet = book.get_sheet_by_index(0)
    rows = sheet.to_python(skip_empty_area=False)
    return [list(row) for row in rows]


def _load_definition_workbook(path: Path):
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        wb = load_workbook(path, data_only=False)
        return wb, wb.active

    if suffix == ".xls":
        rows = _load_calamine_rows(path)
        wb = Workbook()
        ws = wb.active
        for r_idx, row in enumerate(rows, start=1):
            for c_idx, value in enumerate(row, start=1):
                ws.cell(r_idx, c_idx, value)
        return wb, ws

    raise ValueError(f"فرمت فایل گزارش تعریف کالا پشتیبانی نمی‌شود: {suffix}")


def _delete_rows_in_blocks(ws, rows_to_delete: list[int]) -> None:
    if not rows_to_delete:
        return
    rows = sorted(set(rows_to_delete))
    blocks: list[tuple[int, int]] = []
    start = prev = rows[0]
    for row in rows[1:]:
        if row == prev + 1:
            prev = row
            continue
        blocks.append((start, prev))
        start = prev = row
    blocks.append((start, prev))

    for start, end in reversed(blocks):
        ws.delete_rows(start, end - start + 1)


def _first_empty_output_columns(ws, header_row: int, needed: int = 2) -> list[int]:
    result = []
    # Prefer genuinely empty columns already present in the sheet.
    max_col = min(ws.max_column, MAX_HEADER_SCAN_COLS)
    max_data_row = min(ws.max_row, header_row + 2000)
    for col in range(1, max_col + 1):
        if ws.cell(header_row, col).value not in (None, ""):
            continue
        if all(ws.cell(row, col).value in (None, "") for row in range(header_row + 1, max_data_row + 1)):
            result.append(col)
            if len(result) == needed:
                return result

    next_col = ws.max_column + 1
    while len(result) < needed:
        result.append(next_col)
        next_col += 1
    return result


def _last_meaningful_row(ws, header_row: int, columns: Iterable[int]) -> int:
    columns = list(columns)
    for row in range(ws.max_row, header_row, -1):
        if any(ws.cell(row, col).value not in (None, "") for col in columns):
            return row
    return header_row


DEFINITION_DATE_ALIASES = ("تاریخ تعریف", "تاريخ تعريف", "تاریخ ایجاد", "تاريخ ايجاد")
BARCODE_ALIASES = ("بارکد", "باركد", "barcode")
SALES_ALIASES = ("تعداد فروش", "تعداد فروش رفته", "فروش تعداد")
INITIAL_STOCK_ALIASES = (
    "موجودی اولیه",
    "موجودي اوليه",
    "موجودی اول دوره",
    "موجودي اول دوره",
    "موجودی ابتدای دوره",
)
ISSUER_ALIASES = ("صادر کننده", "صادرکننده", "صادر كننده", "صادركننده")
TECHNICAL_ALIASES = ("مشخصات فنی", "مشخصات فني", "مشخصه فنی", "مشخصه فني")
DOCUMENT_DATE_ALIASES = ("تاریخ سند", "تاريخ سند")



def _find_initial_stock_col_ws(ws, header_row: int) -> int:
    """Find ONLY the initial-stock column; never confuse it with plain «موجودی»."""
    max_col = min(ws.max_column, MAX_HEADER_SCAN_COLS)

    # Important: matching is one-way.  «موجودی» must NOT match
    # «موجودی اولیه» merely because it is a substring of the alias.
    for col in range(1, max_col + 1):
        value = ws.cell(header_row, col).value
        value_norm = _norm_text(value)
        value_compact = _compact_text(value)
        if not value_norm:
            continue

        for alias in INITIAL_STOCK_ALIASES:
            alias_norm = _norm_text(alias)
            alias_compact = _compact_text(alias)

            if value_norm == alias_norm or value_compact == alias_compact:
                return col

            # Allow harmless suffixes such as «موجودی اولیه کالا», while still
            # refusing the shorter/current-stock header «موجودی».
            if value_norm.startswith(alias_norm + " "):
                return col
            if value_compact.startswith(alias_compact) and len(value_compact) > len(alias_compact):
                return col

    visible = [
        str(ws.cell(header_row, col).value)
        for col in range(1, max_col + 1)
        if ws.cell(header_row, col).value not in (None, "")
    ]
    raise ValueError(
        "ستون «موجودی اولیه» پیدا نشد. برای محاسبه درصد فروش، "
        "ستون «موجودی» به‌جای آن استفاده نمی‌شود. ستون‌های موجود: "
        + ", ".join(visible[:60])
    )



# Final Excel visual style requested by the user.
PEYDA_FONT_NAME = "Peyda"
WHITE_FILL = PatternFill(fill_type="solid", fgColor="FFFFFFFF")
HEADER_FILL = PatternFill(fill_type="solid", fgColor="FFD9D9D9")
GRID_SIDE = Side(style="thin", color="FFD0D0D0")
GRID_BORDER = Border(left=GRID_SIDE, right=GRID_SIDE, top=GRID_SIDE, bottom=GRID_SIDE)


def _cell_has_duplicate_red_fill(cell) -> bool:
    """Keep the red similarity markers when normalizing the rest of the sheet."""
    if cell.fill is None or cell.fill.fill_type != "solid":
        return False
    color = cell.fill.fgColor
    rgb = getattr(color, "rgb", None)
    if isinstance(rgb, str):
        return rgb.upper() in {"FFFF0000", "FFFFC7CE", "00FF0000", "00FFC7CE"}
    return False


def _apply_final_output_style(ws, header_row: int) -> None:
    """Apply Peyda + clean white/gray styling to the final used table."""
    max_row = ws.max_row
    max_col = ws.max_column

    # Persian workbook usability.
    try:
        ws.sheet_view.rightToLeft = True
    except Exception:
        pass

    # Header row: gray, bold, black, Peyda.
    for col in range(1, max_col + 1):
        cell = ws.cell(header_row, col)
        cell.font = Font(
            name=PEYDA_FONT_NAME,
            size=11,
            bold=True,
            color="FF000000",
        )
        cell.fill = HEADER_FILL
        cell.border = GRID_BORDER
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    ws.row_dimensions[header_row].height = 26

    # Data area: white background + black Peyda font.
    # The two matched cells (barcode/technical) keep their red fill.
    for row in range(header_row + 1, max_row + 1):
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 0, 20)
        for col in range(1, max_col + 1):
            cell = ws.cell(row, col)
            keep_red = _cell_has_duplicate_red_fill(cell)
            cell.font = Font(
                name=PEYDA_FONT_NAME,
                size=10.5,
                bold=False,
                color="FF000000",
            )
            if keep_red:
                cell.fill = RED_FILL
            else:
                cell.fill = WHITE_FILL
            cell.border = GRID_BORDER
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )


def create_sales_percentage_report(definition_path, transfer_path, output_path):
    definition_path = Path(definition_path)
    transfer_path = Path(transfer_path)
    output_path = Path(output_path)

    try:
        wb_definition, ws_definition = _load_definition_workbook(definition_path)
        transfer_rows = _load_calamine_rows(transfer_path)
        if not transfer_rows:
            raise ValueError("فایل انتقال خالی است.")

        definition_header = _find_header_row_ws(
            ws_definition,
            [DEFINITION_DATE_ALIASES, BARCODE_ALIASES, SALES_ALIASES, INITIAL_STOCK_ALIASES],
        )
        transfer_header_idx = _find_header_row_matrix(
            transfer_rows,
            [ISSUER_ALIASES, TECHNICAL_ALIASES, DOCUMENT_DATE_ALIASES],
        )

        definition_date_col = _find_col_ws(ws_definition, definition_header, DEFINITION_DATE_ALIASES)
        barcode_col = _find_col_ws(ws_definition, definition_header, BARCODE_ALIASES)
        sales_col = _find_col_ws(ws_definition, definition_header, SALES_ALIASES)
        initial_stock_col = _find_initial_stock_col_ws(ws_definition, definition_header)

        transfer_header = transfer_rows[transfer_header_idx]
        issuer_idx = _find_col_matrix(transfer_header, ISSUER_ALIASES)
        technical_idx = _find_col_matrix(transfer_header, TECHNICAL_ALIASES)
        document_date_idx = _find_col_matrix(transfer_header, DOCUMENT_DATE_ALIASES)

        kept_transfer: list[tuple[object, object]] = []
        removed_issuer_rows = 0
        document_dates: list[tuple[int, int, int]] = []

        for row in transfer_rows[transfer_header_idx + 1 :]:
            issuer = _get_row_value(row, issuer_idx)
            technical = _get_row_value(row, technical_idx)
            document_date_value = _get_row_value(row, document_date_idx)

            if all(v in (None, "") for v in (issuer, technical, document_date_value)):
                continue
            if _issuer_is_blocked(issuer):
                removed_issuer_rows += 1
                continue

            kept_transfer.append((technical, document_date_value))
            parsed_date = _parse_jalali(document_date_value)
            if parsed_date:
                document_dates.append(parsed_date)

        if not kept_transfer:
            raise ValueError("بعد از حذف صادرکننده‌های مشخص‌شده، هیچ ردیفی در فایل انتقال باقی نماند.")
        if not document_dates:
            samples = [str(v) for _, v in kept_transfer[:8]]
            raise ValueError(
                "در ستون «تاریخ سند» تاریخ معتبر پیدا نشد. چند مقدار اول: " + " | ".join(samples)
            )

        document_date = Counter(document_dates).most_common(1)[0][0]
        next_date = _jalali_next_day(document_date)
        allowed_dates = {document_date, next_date}

        last_definition_row = _last_meaningful_row(
            ws_definition,
            definition_header,
            [definition_date_col, barcode_col, sales_col, initial_stock_col],
        )

        rows_to_delete = []
        for row in range(definition_header + 1, last_definition_row + 1):
            parsed = _parse_jalali(ws_definition.cell(row, definition_date_col).value)
            if parsed not in allowed_dates:
                rows_to_delete.append(row)
        _delete_rows_in_blocks(ws_definition, rows_to_delete)

        # Resolve columns again after row deletion.
        barcode_col = _find_col_ws(ws_definition, definition_header, BARCODE_ALIASES)
        sales_col = _find_col_ws(ws_definition, definition_header, SALES_ALIASES)
        initial_stock_col = _find_initial_stock_col_ws(ws_definition, definition_header)

        # Always place «درصد فروش» directly after «تعداد فروش».
        existing_percentage_col = _find_col_ws(
            ws_definition,
            definition_header,
            ("درصد فروش",),
            required=False,
        )
        if existing_percentage_col is not None:
            ws_definition.delete_cols(existing_percentage_col, 1)
            sales_col = _find_col_ws(ws_definition, definition_header, SALES_ALIASES)
            initial_stock_col = _find_initial_stock_col_ws(ws_definition, definition_header)
            barcode_col = _find_col_ws(ws_definition, definition_header, BARCODE_ALIASES)

        percentage_col = sales_col + 1
        ws_definition.insert_cols(percentage_col, 1)
        _copy_cell_style(
            ws_definition.cell(definition_header, sales_col),
            ws_definition.cell(definition_header, percentage_col),
        )
        ws_definition.cell(definition_header, percentage_col, "درصد فروش")

        sales_col = _find_col_ws(ws_definition, definition_header, SALES_ALIASES)
        initial_stock_col = _find_initial_stock_col_ws(ws_definition, definition_header)
        barcode_col = _find_col_ws(ws_definition, definition_header, BARCODE_ALIASES)
        percentage_col = _find_col_ws(ws_definition, definition_header, ("درصد فروش",))

        last_definition_row = _last_meaningful_row(
            ws_definition,
            definition_header,
            [barcode_col, sales_col, initial_stock_col],
        )

        calculated_percentage_rows = 0
        for row in range(definition_header + 1, last_definition_row + 1):
            sales_value = _to_number(ws_definition.cell(row, sales_col).value)
            stock_value = _to_number(ws_definition.cell(row, initial_stock_col).value)
            if sales_value is None or stock_value is None:
                percentage = None
            elif stock_value == 0:
                percentage = 0
            else:
                raw_percentage = (sales_value / stock_value) * 100
                # کاربر خروجی ساده می‌خواهد: عدد صحیح، بدون اعشار و بدون علامت درصد.
                percentage = (
                    math.floor(raw_percentage + 0.5)
                    if raw_percentage >= 0
                    else math.ceil(raw_percentage - 0.5)
                )
                calculated_percentage_rows += 1
            cell = ws_definition.cell(row, percentage_col, percentage)
            _copy_cell_style(ws_definition.cell(row, sales_col), cell)
            cell.number_format = "0"

        # Put transfer columns into available empty columns of the first workbook.
        technical_output_col, document_date_output_col = _first_empty_output_columns(
            ws_definition, definition_header, 2
        )
        ws_definition.cell(definition_header, technical_output_col, "مشخصات فنی")
        ws_definition.cell(definition_header, document_date_output_col, "تاریخ سند")
        _copy_cell_style(
            ws_definition.cell(definition_header, 1),
            ws_definition.cell(definition_header, technical_output_col),
        )
        _copy_cell_style(
            ws_definition.cell(definition_header, 1),
            ws_definition.cell(definition_header, document_date_output_col),
        )

        # Index transfer rows by technical specification. This lets the final output
        # keep only true barcode/technical matches and place the matched values on
        # the same surviving row instead of leaving unrelated transfer rows beside them.
        transfer_matches: dict[str, list[tuple[object, object]]] = {}
        for technical, date_value in kept_transfer:
            key = _norm_key(technical)
            if key:
                transfer_matches.setdefault(key, []).append((technical, date_value))

        definition_keys: dict[str, list[int]] = {}
        last_definition_row = _last_meaningful_row(
            ws_definition,
            definition_header,
            [barcode_col],
        )
        for row in range(definition_header + 1, last_definition_row + 1):
            key = _norm_key(ws_definition.cell(row, barcode_col).value)
            if key:
                definition_keys.setdefault(key, []).append(row)

        duplicate_keys = set(transfer_matches).intersection(definition_keys)

        # First mark all matched barcode/technical pairs red, as requested.
        for key in duplicate_keys:
            technical_value, date_value = transfer_matches[key][0]
            for row in definition_keys[key]:
                ws_definition.cell(row, technical_output_col, technical_value)
                ws_definition.cell(row, document_date_output_col, date_value)
                ws_definition.cell(row, barcode_col).fill = RED_FILL
                ws_definition.cell(row, technical_output_col).fill = RED_FILL

        # Then remove every data row whose barcode was not one of the red/matched codes.
        # The final workbook therefore contains only similarities between «بارکد» and
        # «مشخصات فنی».
        unmatched_rows = []
        for row in range(definition_header + 1, last_definition_row + 1):
            key = _norm_key(ws_definition.cell(row, barcode_col).value)
            if key not in duplicate_keys:
                unmatched_rows.append(row)
        _delete_rows_in_blocks(ws_definition, unmatched_rows)

        # Resolve columns/last row again because whole rows were deleted.
        barcode_col = _find_col_ws(ws_definition, definition_header, BARCODE_ALIASES)
        sales_col = _find_col_ws(ws_definition, definition_header, SALES_ALIASES)
        initial_stock_col = _find_initial_stock_col_ws(ws_definition, definition_header)
        percentage_col = _find_col_ws(ws_definition, definition_header, ("درصد فروش",))
        technical_output_col = _find_col_ws(ws_definition, definition_header, TECHNICAL_ALIASES)
        document_date_output_col = _find_col_ws(ws_definition, definition_header, DOCUMENT_DATE_ALIASES)
        last_definition_row = _last_meaningful_row(
            ws_definition, definition_header, [barcode_col]
        )

        # Remove requested columns from final output. Do it right-to-left.
        remove_cols = []
        for name in ("شناسه راهکاران", "قیمت", "تاریخ ویرایش"):
            col = _find_col_ws(ws_definition, definition_header, (name,), required=False)
            if col is not None:
                remove_cols.append(col)
        for col in sorted(set(remove_cols), reverse=True):
            ws_definition.delete_cols(col, 1)

        percentage_col = _find_col_ws(ws_definition, definition_header, ("درصد فروش",), required=False)
        technical_output_col = _find_col_ws(ws_definition, definition_header, TECHNICAL_ALIASES, required=False)
        document_date_output_col = _find_col_ws(ws_definition, definition_header, DOCUMENT_DATE_ALIASES, required=False)

        if percentage_col:
            letter = get_column_letter(percentage_col)
            ws_definition.column_dimensions[letter].width = max(
                ws_definition.column_dimensions[letter].width or 0, 13
            )
        if technical_output_col:
            ws_definition.column_dimensions[get_column_letter(technical_output_col)].width = 22
        if document_date_output_col:
            ws_definition.column_dimensions[get_column_letter(document_date_output_col)].width = 16

        # Normalize the final workbook appearance after all filtering/deletions.
        _apply_final_output_style(ws_definition, definition_header)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb_definition.save(output_path)

        year, month, day = document_date
        return {
            "row_count": max(0, last_definition_row - definition_header),
            "removed_issuer_count": removed_issuer_rows,
            "duplicate_count": len(duplicate_keys),
            "percentage_count": calculated_percentage_rows,
            "document_date": f"{year:04d}/{month:02d}/{day:02d}",
            "next_date": f"{next_date[0]:04d}/{next_date[1]:02d}/{next_date[2]:02d}",
        }

    except Exception:
        log_path = Path(__file__).resolve().parent.parent / "sales_percentage_error.log"
        try:
            with log_path.open("a", encoding="utf-8") as log:
                log.write("\n" + "=" * 90 + "\n")
                log.write(f"definition: {definition_path} ({definition_path.stat().st_size if definition_path.exists() else 'missing'} bytes)\n")
                log.write(f"transfer:   {transfer_path} ({transfer_path.stat().st_size if transfer_path.exists() else 'missing'} bytes)\n")
                log.write(traceback.format_exc())
        except Exception:
            pass
        raise
