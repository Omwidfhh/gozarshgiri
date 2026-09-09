from copy import copy
from math import isfinite
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter


PERSIAN_DIGITS = str.maketrans(
    "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
    "01234567890123456789",
)

RULE_COLORS = {
    "soft_red": "F4CCCC",
    "soft_orange": "FCE5CD",
    "soft_green": "D9EAD3",
    "soft_blue": "D9EAF7",
    "soft_yellow": "FFF2CC",
}

MAX_AUTOFIT_WIDTH = 45
MIN_AUTOFIT_WIDTH = 10


def has_value(value):
    if value is None:
        return False

    if isinstance(value, str):
        return bool(value.strip())

    return True


def normalize_text(value):
    if value is None:
        return ""

    return (
        str(value)
        .translate(PERSIAN_DIGITS)
        .replace("ي", "ی")
        .replace("ك", "ک")
        .replace("\u200c", " ")
        .replace("\u200e", "")
        .replace("\u200f", "")
        .replace("\xa0", " ")
        .strip()
    )


def normalize_header(value):
    return "".join(
        normalize_text(value).split()
    ).casefold()


def normalize_code(value):
    if value is None:
        return ""

    if isinstance(value, bool):
        return str(value)

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float):
        if not isfinite(value):
            return ""

        if value.is_integer():
            return str(int(value))

        return format(value, ".15g")

    return "".join(
        normalize_text(value).split()
    )


def display_header(value, column_number):
    text = normalize_text(value)

    if text:
        return text

    return f"ستون {get_column_letter(column_number)}"


def has_valid_dimensions(worksheet):
    return (
        isinstance(worksheet.max_row, int)
        and worksheet.max_row >= 1
        and isinstance(worksheet.max_column, int)
        and worksheet.max_column >= 1
    )


def load_compatible_workbook(
    file_path,
    *,
    data_only,
    prefer_read_only,
):
    workbook = load_workbook(
        filename=file_path,
        read_only=prefer_read_only,
        data_only=data_only,
    )
    worksheet = workbook.active

    if (
        not prefer_read_only
        or has_valid_dimensions(worksheet)
    ):
        return workbook, worksheet

    # بعضی خروجی‌های نرم‌افزارهای حسابداری تگ dimension ندارند.
    # در حالت read_only ابعاد چنین شیتی None می‌شود؛ حالت عادی
    # ابعاد را مستقیماً از سلول‌ها محاسبه می‌کند.
    workbook.close()
    workbook = load_workbook(
        filename=file_path,
        read_only=False,
        data_only=data_only,
    )

    return workbook, workbook.active


def get_headers(worksheet):
    last_column = worksheet.max_column or 1

    return [
        display_header(
            worksheet.cell(
                row=1,
                column=column_number,
            ).value,
            column_number,
        )
        for column_number in range(1, last_column + 1)
    ]


def make_unique_headers(headers):
    counts = {}
    unique_headers = []

    for header in headers:
        normalized = normalize_header(header)
        count = counts.get(normalized, 0) + 1
        counts[normalized] = count

        if count == 1:
            unique_headers.append(header)
        else:
            unique_headers.append(
                f"{header} ({count})"
            )

    return unique_headers


def inspect_excel_file(file_path, file_label):
    workbook, worksheet = load_compatible_workbook(
        file_path,
        data_only=True,
        prefer_read_only=True,
    )

    try:
        raw_headers = get_headers(worksheet)
        headers = make_unique_headers(raw_headers)

        return {
            "label": file_label,
            "sheet_name": worksheet.title,
            "row_count": max(
                (worksheet.max_row or 1) - 1,
                0,
            ),
            "column_count": len(headers),
            "headers": [
                {
                    "name": header,
                    "source_name": raw_headers[index],
                    "index": index + 1,
                    "letter": get_column_letter(index + 1),
                }
                for index, header in enumerate(headers)
            ],
        }

    finally:
        workbook.close()


def resolve_header_index(headers, requested_header, file_label):
    requested_normalized = normalize_header(
        requested_header
    )

    for index, header in enumerate(headers, start=1):
        if normalize_header(header) == requested_normalized:
            return index

    raise ValueError(
        f"ستون «{requested_header}» در {file_label} پیدا نشد."
    )


def safe_output_header(requested_name, existing_headers):
    base_name = normalize_text(requested_name) or "ستون جدید"
    existing_normalized = {
        normalize_header(header)
        for header in existing_headers
    }

    if normalize_header(base_name) not in existing_normalized:
        return base_name

    counter = 2

    while True:
        candidate = f"{base_name} ({counter})"

        if normalize_header(candidate) not in existing_normalized:
            return candidate

        counter += 1


def copy_cell_style(source_cell, target_cell):
    if source_cell.has_style:
        target_cell._style = copy(source_cell._style)

    if source_cell.number_format:
        target_cell.number_format = source_cell.number_format

    target_cell.alignment = copy(source_cell.alignment)
    target_cell.border = copy(source_cell.border)
    target_cell.fill = copy(source_cell.fill)
    target_cell.protection = copy(source_cell.protection)


def build_lookup_map(
    worksheet,
    key_column,
    append_columns,
):
    lookup_map = {}

    for row_number in range(2, worksheet.max_row + 1):
        key = normalize_code(
            worksheet.cell(
                row=row_number,
                column=key_column,
            ).value
        )

        if not key:
            continue

        values = {
            item["output_header"]: worksheet.cell(
                row=row_number,
                column=item["source_column"],
            ).value
            for item in append_columns
        }

        if key not in lookup_map:
            lookup_map[key] = values
            continue

        for output_header, value in values.items():
            if (
                not has_value(lookup_map[key].get(output_header))
                and has_value(value)
            ):
                lookup_map[key][output_header] = value

    return lookup_map


def coerce_number(value):
    if value is None or isinstance(value, bool):
        return None

    if isinstance(value, (int, float)):
        return float(value)

    text = normalize_text(value).replace(",", "")

    try:
        return float(text)
    except (TypeError, ValueError):
        return None


def values_equal(left, right):
    left_number = coerce_number(left)
    right_number = coerce_number(right)

    if left_number is not None and right_number is not None:
        return left_number == right_number

    return normalize_text(left).casefold() == normalize_text(right).casefold()


def rule_matches(value, operator, compare_value):
    if operator == "empty":
        return not has_value(value)

    if operator == "not_empty":
        return has_value(value)

    if operator == "contains":
        return (
            normalize_text(compare_value).casefold()
            in normalize_text(value).casefold()
        )

    if operator == "not_contains":
        return (
            normalize_text(compare_value).casefold()
            not in normalize_text(value).casefold()
        )

    if operator == "eq":
        return values_equal(value, compare_value)

    if operator == "neq":
        return not values_equal(value, compare_value)

    value_number = coerce_number(value)
    compare_number = coerce_number(compare_value)

    if value_number is None or compare_number is None:
        return False

    if operator == "gt":
        return value_number > compare_number

    if operator == "gte":
        return value_number >= compare_number

    if operator == "lt":
        return value_number < compare_number

    if operator == "lte":
        return value_number <= compare_number

    return False


def apply_highlight_rules(
    worksheet,
    rules,
):
    if not rules or worksheet.max_row < 2:
        return

    headers = get_headers(worksheet)
    prepared_rules = []

    for rule in rules:
        column_name = rule.get("column")
        operator = rule.get("operator")
        color_name = rule.get("color", "soft_red")

        if not column_name or not operator:
            continue

        column_number = resolve_header_index(
            headers,
            column_name,
            "خروجی نهایی",
        )
        color = RULE_COLORS.get(
            color_name,
            RULE_COLORS["soft_red"],
        )

        prepared_rules.append({
            "column": column_number,
            "operator": operator,
            "value": rule.get("value"),
            "fill": PatternFill(
                fill_type="solid",
                fgColor=color,
            ),
        })

    for row_number in range(2, worksheet.max_row + 1):
        for rule in prepared_rules:
            value = worksheet.cell(
                row=row_number,
                column=rule["column"],
            ).value

            if not rule_matches(
                value,
                rule["operator"],
                rule["value"],
            ):
                continue

            for column_number in range(
                1,
                worksheet.max_column + 1,
            ):
                worksheet.cell(
                    row=row_number,
                    column=column_number,
                ).fill = copy(rule["fill"])

            break


def apply_peyda_font(worksheet):
    for row in worksheet.iter_rows():
        for cell in row:
            old_font = cell.font
            cell.font = Font(
                name="Peyda",
                size=11,
                bold=old_font.bold,
                italic=old_font.italic,
                color=copy(old_font.color),
                underline=old_font.underline,
                strike=old_font.strike,
            )


def autofit_columns(worksheet):
    for column_number in range(
        1,
        worksheet.max_column + 1,
    ):
        maximum_length = 0

        for row_number in range(
            1,
            worksheet.max_row + 1,
        ):
            value = worksheet.cell(
                row=row_number,
                column=column_number,
            ).value

            if value is None:
                continue

            maximum_length = max(
                maximum_length,
                len(str(value)),
            )

        worksheet.column_dimensions[
            get_column_letter(column_number)
        ].width = min(
            max(maximum_length + 3, MIN_AUTOFIT_WIDTH),
            MAX_AUTOFIT_WIDTH,
        )


def validate_config(config):
    required_fields = (
        "name",
        "base_key",
        "lookup_key",
    )

    for field in required_fields:
        if not normalize_text(config.get(field)):
            raise ValueError(
                "تنظیمات گزارش‌ساز کامل نیست."
            )

    append_columns = config.get("append_columns") or []

    if not append_columns:
        raise ValueError(
            "حداقل یک ستون برای انتقال از اکسل دوم انتخاب کنید."
        )

    if len(append_columns) > 20:
        raise ValueError(
            "حداکثر ۲۰ ستون را می‌توانید منتقل کنید."
        )

    if len(config.get("rules") or []) > 5:
        raise ValueError(
            "حداکثر ۵ قانون رنگ مجاز است."
        )


def create_custom_report(
    base_file_path,
    lookup_file_path,
    config,
    output_path,
):
    validate_config(config)

    base_file_path = Path(base_file_path)
    lookup_file_path = Path(lookup_file_path)
    output_path = Path(output_path)
    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    base_workbook, base_sheet = load_compatible_workbook(
        base_file_path,
        data_only=False,
        prefer_read_only=False,
    )
    lookup_workbook, lookup_sheet = load_compatible_workbook(
        lookup_file_path,
        data_only=True,
        prefer_read_only=True,
    )

    try:
        base_headers = get_headers(base_sheet)
        lookup_headers = get_headers(lookup_sheet)

        base_key_column = resolve_header_index(
            base_headers,
            config["base_key"],
            "اکسل مبنا",
        )
        lookup_key_column = resolve_header_index(
            lookup_headers,
            config["lookup_key"],
            "اکسل دوم",
        )

        removed_headers = config.get("remove_columns") or []
        removed_normalized = {
            normalize_header(header)
            for header in removed_headers
        }

        if normalize_header(config["base_key"]) in removed_normalized:
            raise ValueError(
                "ستون تطبیق اکسل مبنا نباید حذف شود."
            )

        output_headers = list(base_headers)
        prepared_append_columns = []
        rule_header_aliases = {}

        for item in config["append_columns"]:
            source_header = item.get("source_header")

            if not normalize_text(source_header):
                continue

            source_column = resolve_header_index(
                lookup_headers,
                source_header,
                "اکسل دوم",
            )
            requested_output = (
                item.get("output_header")
                or source_header
            )
            output_header = safe_output_header(
                requested_output,
                output_headers,
            )
            output_headers.append(output_header)
            rule_header_aliases[
                normalize_header(requested_output)
            ] = output_header

            prepared_append_columns.append({
                "source_column": source_column,
                "source_header": source_header,
                "output_header": output_header,
            })

        if not prepared_append_columns:
            raise ValueError(
                "ستون معتبری برای انتقال انتخاب نشده است."
            )

        lookup_map = build_lookup_map(
            lookup_sheet,
            lookup_key_column,
            prepared_append_columns,
        )

        original_last_column = base_sheet.max_column
        header_style_source = base_sheet.cell(
            row=1,
            column=original_last_column,
        )

        for offset, item in enumerate(
            prepared_append_columns,
            start=1,
        ):
            column_number = original_last_column + offset
            header_cell = base_sheet.cell(
                row=1,
                column=column_number,
            )
            copy_cell_style(
                header_style_source,
                header_cell,
            )
            header_cell.value = item["output_header"]

        matched_count = 0
        unmatched_count = 0

        for row_number in range(2, base_sheet.max_row + 1):
            key = normalize_code(
                base_sheet.cell(
                    row=row_number,
                    column=base_key_column,
                ).value
            )
            lookup_values = lookup_map.get(key)

            if key and lookup_values is not None:
                matched_count += 1
            elif key:
                unmatched_count += 1

            style_source = base_sheet.cell(
                row=row_number,
                column=original_last_column,
            )

            for offset, item in enumerate(
                prepared_append_columns,
                start=1,
            ):
                target_cell = base_sheet.cell(
                    row=row_number,
                    column=original_last_column + offset,
                )
                copy_cell_style(
                    style_source,
                    target_cell,
                )
                target_cell.value = (
                    lookup_values.get(item["output_header"])
                    if lookup_values is not None
                    else None
                )

        columns_to_delete = []

        for column_number, header in enumerate(
            base_headers,
            start=1,
        ):
            if normalize_header(header) in removed_normalized:
                columns_to_delete.append(column_number)

        for column_number in sorted(
            columns_to_delete,
            reverse=True,
        ):
            base_sheet.delete_cols(
                column_number,
                1,
            )

        resolved_rules = []

        for rule in config.get("rules") or []:
            resolved_rule = dict(rule)
            normalized_rule_column = normalize_header(
                rule.get("column")
            )

            if normalized_rule_column in rule_header_aliases:
                resolved_rule["column"] = rule_header_aliases[
                    normalized_rule_column
                ]

            resolved_rules.append(resolved_rule)

        apply_highlight_rules(
            base_sheet,
            resolved_rules,
        )
        apply_peyda_font(base_sheet)
        autofit_columns(base_sheet)

        base_sheet.freeze_panes = "A2"

        if base_sheet.max_row >= 1:
            base_sheet.auto_filter.ref = (
                f"A1:"
                f"{get_column_letter(base_sheet.max_column)}"
                f"{max(base_sheet.max_row, 1)}"
            )

        base_sheet.sheet_view.rightToLeft = True
        base_workbook.save(output_path)

        return {
            "output_path": str(output_path),
            "matched_count": matched_count,
            "unmatched_count": unmatched_count,
            "row_count": max(base_sheet.max_row - 1, 0),
        }

    finally:
        lookup_workbook.close()
        base_workbook.close()
