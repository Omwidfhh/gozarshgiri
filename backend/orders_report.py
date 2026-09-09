from copy import copy
from decimal import Decimal
from decimal import InvalidOperation
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.styles import Border
from openpyxl.styles import PatternFill
from openpyxl.styles import Side
from openpyxl.utils import get_column_letter


ADMIN_ORDER_HEADERS = (
    "شماره راهکاران",
)

ADMIN_AMOUNT_HEADERS = (
    "مبلغ",
)

RAHKARAN_ORDER_HEADERS = (
    "شماره سفارش",
)

RAHKARAN_NET_AMOUNT_HEADERS = (
    "مبلغ خالص",
)

NET_AMOUNT_OUTPUT_HEADER = "مبلغ خالص"
DIFFERENCE_OUTPUT_HEADER = "اختلاف مبلغ"
STATUS_OUTPUT_HEADER = "وضعیت"

PEYDA_FONT_NAME = "Peyda"
PEYDA_FONT_SIZE = 11
AMOUNT_NUMBER_FORMAT = "#,##0"

LIGHT_RED_FILL = PatternFill(
    fill_type="solid",
    fgColor="FCE8E6",
)

THIN_GRAY_SIDE = Side(
    style="thin",
    color="D9D9D9",
)

ALL_BORDERS = Border(
    left=THIN_GRAY_SIDE,
    right=THIN_GRAY_SIDE,
    top=THIN_GRAY_SIDE,
    bottom=THIN_GRAY_SIDE,
)


PERSIAN_DIGITS = str.maketrans(
    "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
    "01234567890123456789",
)


def has_value(value):
    return value is not None and str(value).strip() != ""


def normalize_text(value):
    if value is None:
        return ""

    return (
        str(value)
        .translate(PERSIAN_DIGITS)
        .replace("ي", "ی")
        .replace("ك", "ک")
        .replace("\u200c", " ")
        .replace("\u200f", "")
        .replace("\u200e", "")
        .replace("\xa0", " ")
        .strip()
    )


def normalize_header(value):
    return "".join(
        normalize_text(value).split()
    )


def normalize_order_id(value):
    if value is None:
        return ""

    if isinstance(value, bool):
        return str(value).strip()

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))

        return format(value, ".15g")

    text = normalize_text(value)

    if re.fullmatch(r"[+-]?\d+\.0+", text):
        return text.split(".")[0]

    return text


def parse_amount(value):
    if value is None:
        return Decimal("0")

    if isinstance(value, bool):
        return Decimal(int(value))

    if isinstance(value, Decimal):
        return value

    if isinstance(value, int):
        return Decimal(value)

    if isinstance(value, float):
        return Decimal(str(value))

    text = normalize_text(value)

    if not text:
        return Decimal("0")

    is_negative = (
        text.startswith("(")
        and text.endswith(")")
    )

    text = (
        text.replace("(", "")
        .replace(")", "")
        .replace(",", "")
        .replace("٬", "")
        .replace("٫", ".")
        .replace(" ", "")
        .replace("\xa0", "")
    )

    number_match = re.search(
        r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?",
        text,
    )

    if not number_match:
        return Decimal("0")

    try:
        amount = Decimal(number_match.group(0))

        if is_negative and amount > 0:
            amount = -amount

        return amount

    except InvalidOperation:
        return Decimal("0")


def excel_number(value):
    if value == value.to_integral_value():
        return int(value)

    return float(value)


def safe_max_row(worksheet):
    max_row = worksheet.max_row

    if max_row is None:
        return 1

    try:
        return max(int(max_row), 1)

    except (TypeError, ValueError):
        return 1


def safe_max_column(worksheet):
    max_column = worksheet.max_column

    if max_column is None:
        return 1

    try:
        return max(int(max_column), 1)

    except (TypeError, ValueError):
        return 1


def find_last_row(
    worksheet,
    column_number,
):
    maximum_row = safe_max_row(worksheet)

    for row_number in range(
        maximum_row,
        0,
        -1,
    ):
        value = worksheet.cell(
            row=row_number,
            column=column_number,
        ).value

        if has_value(value):
            return row_number

    return 1


def find_last_column(
    worksheet,
    row_number=1,
):
    maximum_column = safe_max_column(worksheet)

    for column_number in range(
        maximum_column,
        0,
        -1,
    ):
        value = worksheet.cell(
            row=row_number,
            column=column_number,
        ).value

        if has_value(value):
            return column_number

    return 1


def find_header_column(
    worksheet,
    accepted_headers,
):
    normalized_headers = {
        normalize_header(header)
        for header in accepted_headers
    }

    last_column = find_last_column(
        worksheet,
        row_number=1,
    )

    for column_number in range(
        1,
        last_column + 1,
    ):
        current_header = normalize_header(
            worksheet.cell(
                row=1,
                column=column_number,
            ).value
        )

        if current_header in normalized_headers:
            return column_number

    return None


def require_header_column(
    worksheet,
    accepted_headers,
    file_label,
):
    column_number = find_header_column(
        worksheet,
        accepted_headers,
    )

    if column_number is None:
        readable_header = accepted_headers[0]

        raise ValueError(
            f"ستون «{readable_header}» "
            f"در {file_label} پیدا نشد."
        )

    return column_number


def remove_old_result_sheet(
    workbook,
    admin_sheet,
):
    for worksheet in list(workbook.worksheets):
        if worksheet is admin_sheet:
            continue

        normalized_name = normalize_header(
            worksheet.title
        )

        if normalized_name == normalize_header("نتیجه"):
            workbook.remove(worksheet)


def copy_admin_sheet(
    workbook,
    admin_sheet,
):
    result_sheet = workbook.copy_worksheet(
        admin_sheet
    )

    result_sheet.title = "نتیجه"

    current_index = workbook.worksheets.index(
        result_sheet
    )

    admin_index = workbook.worksheets.index(
        admin_sheet
    )

    target_index = admin_index + 1

    workbook.move_sheet(
        result_sheet,
        offset=target_index - current_index,
    )

    workbook.active = workbook.worksheets.index(
        result_sheet
    )

    return result_sheet


def write_output_header(
    worksheet,
    column_number,
    title,
    template_column,
):
    template_cell = worksheet.cell(
        row=1,
        column=template_column,
    )

    output_cell = worksheet.cell(
        row=1,
        column=column_number,
        value=title,
    )

    if template_cell.has_style:
        output_cell._style = copy(
            template_cell._style
        )

    output_cell.alignment = copy(
        template_cell.alignment
    )


def calculate_column_width(value):
    if value is None:
        return 0

    text = str(value)
    width = 0

    for character in text:
        if "\u0600" <= character <= "\u06ff":
            width += 1.35
        else:
            width += 1

    return width


def auto_fit_columns(
    worksheet,
    amount_columns=None,
):
    amount_columns = set(
        amount_columns or []
    )

    maximum_row = safe_max_row(worksheet)
    maximum_column = safe_max_column(worksheet)
    maximum_widths = [0] * maximum_column

    for row in worksheet.iter_rows(
        min_row=1,
        max_row=maximum_row,
        min_col=1,
        max_col=maximum_column,
    ):
        for cell in row:
            value_width = calculate_column_width(
                cell.value
            )

            width_index = cell.column - 1

            if value_width > maximum_widths[width_index]:
                maximum_widths[width_index] = value_width

    for column_number, maximum_width in enumerate(
        maximum_widths,
        start=1,
    ):
        column_letter = get_column_letter(
            column_number
        )

        minimum_width = 16 if (
            column_number in amount_columns
        ) else 10

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(
                maximum_width + 3,
                minimum_width,
            ),
            60,
        )

        worksheet.column_dimensions[
            column_letter
        ].bestFit = True


def apply_final_formatting(
    worksheet,
    amount_columns,
    status_column,
):
    maximum_row = safe_max_row(worksheet)
    maximum_column = safe_max_column(worksheet)
    amount_columns = set(amount_columns)

    for row_number in range(
        1,
        maximum_row + 1,
    ):
        status_value = normalize_text(
            worksheet.cell(
                row=row_number,
                column=status_column,
            ).value
        )

        has_difference = (
            row_number > 1
            and status_value == "اختلاف دارد"
        )

        for column_number in range(
            1,
            maximum_column + 1,
        ):
            cell = worksheet.cell(
                row=row_number,
                column=column_number,
            )

            peyda_font = copy(cell.font)
            peyda_font.name = PEYDA_FONT_NAME
            peyda_font.sz = PEYDA_FONT_SIZE

            cell.font = peyda_font
            cell.border = ALL_BORDERS

            if column_number in amount_columns:
                if (
                    row_number > 1
                    and cell.data_type != "f"
                    and isinstance(cell.value, str)
                    and re.search(
                        r"\d",
                        normalize_text(cell.value),
                    )
                ):
                    cell.value = excel_number(
                        parse_amount(cell.value)
                    )

                cell.number_format = (
                    AMOUNT_NUMBER_FORMAT
                )

            if has_difference:
                cell.fill = LIGHT_RED_FILL

    worksheet.sheet_view.rightToLeft = True


def set_amount_cell(
    worksheet,
    row_number,
    column_number,
    amount,
    number_format,
):
    cell = worksheet.cell(
        row=row_number,
        column=column_number,
        value=excel_number(amount),
    )

    if number_format:
        cell.number_format = number_format


def create_order_discrepancy_report(
    admin_file_path,
    rahkaran_file_path,
    output_file_path,
):
    admin_path = Path(admin_file_path)
    rahkaran_path = Path(rahkaran_file_path)
    output_path = Path(output_file_path)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    admin_keep_vba = (
        admin_path.suffix.lower() == ".xlsm"
    )

    rahkaran_keep_vba = (
        rahkaran_path.suffix.lower() == ".xlsm"
    )

    admin_workbook = load_workbook(
        filename=admin_path,
        data_only=False,
        read_only=False,
        keep_vba=admin_keep_vba,
    )

    admin_values_workbook = load_workbook(
        filename=admin_path,
        data_only=True,
        read_only=False,
        keep_vba=admin_keep_vba,
    )

    rahkaran_workbook = load_workbook(
        filename=rahkaran_path,
        data_only=True,
        read_only=False,
        keep_vba=rahkaran_keep_vba,
    )

    try:
        if not admin_workbook.worksheets:
            raise ValueError(
                "فایل ادمین هیچ شیتی ندارد."
            )

        if not rahkaran_workbook.worksheets:
            raise ValueError(
                "فایل راهکاران هیچ شیتی ندارد."
            )

        admin_sheet = admin_workbook.worksheets[0]
        admin_values_sheet = (
            admin_values_workbook.worksheets[0]
        )
        rahkaran_sheet = (
            rahkaran_workbook.worksheets[0]
        )

        admin_order_column = require_header_column(
            worksheet=admin_sheet,
            accepted_headers=ADMIN_ORDER_HEADERS,
            file_label="فایل ادمین",
        )

        admin_amount_column = require_header_column(
            worksheet=admin_sheet,
            accepted_headers=ADMIN_AMOUNT_HEADERS,
            file_label="فایل ادمین",
        )

        rahkaran_order_column = require_header_column(
            worksheet=rahkaran_sheet,
            accepted_headers=RAHKARAN_ORDER_HEADERS,
            file_label="فایل راهکاران",
        )

        rahkaran_net_amount_column = require_header_column(
            worksheet=rahkaran_sheet,
            accepted_headers=RAHKARAN_NET_AMOUNT_HEADERS,
            file_label="فایل راهکاران",
        )

        remove_old_result_sheet(
            workbook=admin_workbook,
            admin_sheet=admin_sheet,
        )

        result_sheet = copy_admin_sheet(
            workbook=admin_workbook,
            admin_sheet=admin_sheet,
        )

        last_admin_column = find_last_column(
            worksheet=result_sheet,
            row_number=1,
        )

        net_amount_output_column = last_admin_column + 1
        difference_output_column = last_admin_column + 2
        status_output_column = last_admin_column + 3

        write_output_header(
            worksheet=result_sheet,
            column_number=net_amount_output_column,
            title=NET_AMOUNT_OUTPUT_HEADER,
            template_column=last_admin_column,
        )

        write_output_header(
            worksheet=result_sheet,
            column_number=difference_output_column,
            title=DIFFERENCE_OUTPUT_HEADER,
            template_column=last_admin_column,
        )

        write_output_header(
            worksheet=result_sheet,
            column_number=status_output_column,
            title=STATUS_OUTPUT_HEADER,
            template_column=last_admin_column,
        )

        amount_number_format = result_sheet.cell(
            row=2,
            column=admin_amount_column,
        ).number_format

        rahkaran_amounts = {}

        last_rahkaran_row = find_last_row(
            worksheet=rahkaran_sheet,
            column_number=rahkaran_order_column,
        )

        for row_number in range(
            2,
            last_rahkaran_row + 1,
        ):
            order_id = normalize_order_id(
                rahkaran_sheet.cell(
                    row=row_number,
                    column=rahkaran_order_column,
                ).value
            )

            if not order_id:
                continue

            net_amount = parse_amount(
                rahkaran_sheet.cell(
                    row=row_number,
                    column=rahkaran_net_amount_column,
                ).value
            )

            # اگر شماره سفارش تکراری باشد،
            # مانند Dictionary در VBA آخرین مقدار استفاده می‌شود.
            rahkaran_amounts[order_id] = net_amount

        admin_orders = set()

        last_admin_row = find_last_row(
            worksheet=result_sheet,
            column_number=admin_order_column,
        )

        for row_number in range(
            2,
            last_admin_row + 1,
        ):
            order_id = normalize_order_id(
                admin_values_sheet.cell(
                    row=row_number,
                    column=admin_order_column,
                ).value
            )

            if not order_id:
                continue

            admin_orders.add(order_id)

            if order_id not in rahkaran_amounts:
                result_sheet.cell(
                    row=row_number,
                    column=status_output_column,
                    value="فقط در ادمین",
                )
                continue

            admin_amount = parse_amount(
                admin_values_sheet.cell(
                    row=row_number,
                    column=admin_amount_column,
                ).value
            )

            net_amount = rahkaran_amounts[order_id]
            difference = net_amount - admin_amount

            set_amount_cell(
                worksheet=result_sheet,
                row_number=row_number,
                column_number=net_amount_output_column,
                amount=net_amount,
                number_format=amount_number_format,
            )

            set_amount_cell(
                worksheet=result_sheet,
                row_number=row_number,
                column_number=difference_output_column,
                amount=difference,
                number_format=amount_number_format,
            )

            if difference == 0:
                status = "اوکی"
            else:
                status = "اختلاف دارد"

            result_sheet.cell(
                row=row_number,
                column=status_output_column,
                value=status,
            )

        result_row = (
            find_last_row(
                worksheet=result_sheet,
                column_number=admin_order_column,
            )
            + 1
        )

        for order_id, net_amount in rahkaran_amounts.items():
            if order_id in admin_orders:
                continue

            result_sheet.cell(
                row=result_row,
                column=admin_order_column,
                value=order_id,
            )

            set_amount_cell(
                worksheet=result_sheet,
                row_number=result_row,
                column_number=net_amount_output_column,
                amount=net_amount,
                number_format=amount_number_format,
            )

            result_sheet.cell(
                row=result_row,
                column=status_output_column,
                value="فقط در راهکاران",
            )

            result_row += 1

        amount_columns = (
            admin_amount_column,
            net_amount_output_column,
            difference_output_column,
        )

        apply_final_formatting(
            worksheet=result_sheet,
            amount_columns=amount_columns,
            status_column=status_output_column,
        )

        auto_fit_columns(
            worksheet=result_sheet,
            amount_columns=amount_columns,
        )

        admin_workbook.save(output_path)

    finally:
        rahkaran_workbook.close()
        admin_values_workbook.close()
        admin_workbook.close()

    return str(output_path)
