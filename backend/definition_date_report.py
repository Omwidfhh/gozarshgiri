from copy import copy
from decimal import Decimal, InvalidOperation
from math import isfinite
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from report import DIGIT_TRANSLATION
from report import PEYDA_BODY_FONT
from report import PEYDA_HEADER_FONT
from report import PRODUCT_STOCK_HEADER
from report import STOCK_HEADER
from report import calculate_visual_width
from report import create_report
from report import find_column_by_aliases
from report import normalize_code


TECHNICAL_ALIASES = (
    "مشخصه فنی",
    "مشخصه فنی کالا",
    "مشخصات فنی",
    "مشخصات فنی کالا",
)

BARCODE_ALIASES = (
    "بارکد",
    "بار کد",
    "کد بارکد",
)

DEFINITION_DATE_ALIASES = (
    "تاریخ تعریف",
    "تاريخ تعريف",
    "تاریخ ثبت",
)

DEFINITION_DATE_HEADER = "تاریخ تعریف"

NO_FILL = PatternFill(
    fill_type=None,
)

LIGHT_RED_FILL = PatternFill(
    fill_type="solid",
    fgColor="F4CCCC",
)

LIGHT_ORANGE_FILL = PatternFill(
    fill_type="solid",
    fgColor="FCE4D6",
)


def get_headers(worksheet, file_label):
    try:
        return next(
            worksheet.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True,
            )
        )
    except StopIteration as error:
        raise ValueError(
            f"{file_label} خالی است."
        ) from error


def find_required_column(
    headers,
    aliases,
    readable_name,
    file_label,
):
    column_index = find_column_by_aliases(
        headers,
        aliases,
    )

    if column_index is None:
        raise ValueError(
            f"ستون «{readable_name}» در {file_label} پیدا نشد."
        )

    return column_index + 1


def parse_stock_number(value):
    if value is None or isinstance(value, bool):
        return None

    if isinstance(value, Decimal):
        return value

    if isinstance(value, int):
        return Decimal(value)

    if isinstance(value, float):
        if not isfinite(value):
            return None

        return Decimal(str(value))

    normalized_value = (
        str(value)
        .translate(DIGIT_TRANSLATION)
        .replace(",", "")
        .replace("٬", "")
        .replace("،", "")
        .strip()
    )

    if not normalized_value:
        return None

    if normalized_value.casefold() in {
        "n/a",
        "na",
        "none",
        "null",
        "-",
    }:
        return None

    is_parenthesized_negative = (
        normalized_value.startswith("(")
        and normalized_value.endswith(")")
    )

    if is_parenthesized_negative:
        normalized_value = (
            f"-{normalized_value[1:-1].strip()}"
        )

    try:
        return Decimal(normalized_value)
    except InvalidOperation:
        return None


def clear_data_row_fills(
    worksheet,
    last_column,
):
    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=last_column,
    ):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue

            cell.fill = copy(NO_FILL)


def fill_entire_row(
    worksheet,
    row_number,
    last_column,
    fill,
):
    for column_number in range(
        1,
        last_column + 1,
    ):
        cell = worksheet.cell(
            row=row_number,
            column=column_number,
        )

        if isinstance(cell, MergedCell):
            continue

        cell.fill = copy(fill)


def apply_inventory_row_colors(
    worksheet,
    site_stock_column,
    product_stock_column,
    last_column,
):
    clear_data_row_fills(
        worksheet=worksheet,
        last_column=last_column,
    )

    zero = Decimal("0")

    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        site_stock = parse_stock_number(
            worksheet.cell(
                row=row_number,
                column=site_stock_column,
            ).value
        )

        product_stock = parse_stock_number(
            worksheet.cell(
                row=row_number,
                column=product_stock_column,
            ).value
        )

        if site_stock is None or product_stock is None:
            continue

        if site_stock != zero and product_stock == zero:
            fill_entire_row(
                worksheet=worksheet,
                row_number=row_number,
                last_column=last_column,
                fill=LIGHT_RED_FILL,
            )
            continue

        if (
            product_stock > zero
            and product_stock < site_stock
        ):
            fill_entire_row(
                worksheet=worksheet,
                row_number=row_number,
                last_column=last_column,
                fill=LIGHT_ORANGE_FILL,
            )


def read_definition_dates(file_path):
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook.active

    try:
        headers = get_headers(
            worksheet,
            "اکسل سوم",
        )

        barcode_column = find_required_column(
            headers,
            BARCODE_ALIASES,
            "بارکد",
            "اکسل سوم",
        )

        definition_date_column = find_required_column(
            headers,
            DEFINITION_DATE_ALIASES,
            "تاریخ تعریف",
            "اکسل سوم",
        )

        required_column = max(
            barcode_column,
            definition_date_column,
        )

        definition_dates = {}

        for row in worksheet.iter_rows(
            min_row=2,
            values_only=False,
        ):
            if len(row) < required_column:
                continue

            barcode = normalize_code(
                row[barcode_column - 1].value
            )

            if not barcode or barcode in definition_dates:
                continue

            date_cell = row[
                definition_date_column - 1
            ]

            definition_dates[barcode] = {
                "value": date_cell.value,
                "number_format": date_cell.number_format,
            }

        return definition_dates

    finally:
        workbook.close()


def append_definition_dates(
    output_path,
    third_file_path,
):
    definition_dates = read_definition_dates(
        third_file_path
    )

    workbook = load_workbook(
        filename=output_path,
        read_only=False,
        data_only=False,
    )

    worksheet = workbook.active

    try:
        headers = get_headers(
            worksheet,
            "خروجی اکسل سایت",
        )

        technical_column = find_required_column(
            headers,
            TECHNICAL_ALIASES,
            "مشخصه فنی",
            "خروجی اکسل سایت",
        )

        stock_column_index = find_column_by_aliases(
            headers,
            (STOCK_HEADER,),
        )

        stock_column = (
            stock_column_index + 1
            if stock_column_index is not None
            else None
        )

        product_stock_column = find_required_column(
            headers,
            (PRODUCT_STOCK_HEADER,),
            PRODUCT_STOCK_HEADER,
            "خروجی اکسل سایت",
        )

        last_header_column = max(
            (
                cell.column
                for cell in worksheet[1]
                if cell.value not in (None, "")
            ),
            default=worksheet.max_column,
        )

        output_column = last_header_column + 1

        source_header_cell = worksheet.cell(
            row=1,
            column=last_header_column,
        )

        output_header_cell = worksheet.cell(
            row=1,
            column=output_column,
        )

        if source_header_cell.has_style:
            output_header_cell._style = copy(
                source_header_cell._style
            )

        output_header_cell.value = (
            DEFINITION_DATE_HEADER
        )
        output_header_cell.font = PEYDA_HEADER_FONT

        maximum_width = calculate_visual_width(
            DEFINITION_DATE_HEADER
        )

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):
            technical_value = normalize_code(
                worksheet.cell(
                    row=row_number,
                    column=technical_column,
                ).value
            )

            matched_date = definition_dates.get(
                technical_value
            )

            adjacent_cell = worksheet.cell(
                row=row_number,
                column=last_header_column,
            )

            output_cell = worksheet.cell(
                row=row_number,
                column=output_column,
            )

            if adjacent_cell.has_style:
                output_cell._style = copy(
                    adjacent_cell._style
                )

            output_cell.font = PEYDA_BODY_FONT

            if matched_date is not None:
                output_cell.value = matched_date[
                    "value"
                ]

                if matched_date["number_format"]:
                    output_cell.number_format = (
                        matched_date["number_format"]
                    )

            maximum_width = max(
                maximum_width,
                calculate_visual_width(
                    output_cell.value
                ),
            )

        output_letter = get_column_letter(
            output_column
        )

        worksheet.column_dimensions[
            output_letter
        ].width = min(
            max(maximum_width + 3, 15),
            30,
        )

        worksheet.column_dimensions[
            output_letter
        ].bestFit = True

        worksheet.auto_filter.ref = (
            f"A1:{output_letter}{worksheet.max_row}"
        )

        if stock_column is None:
            raise ValueError(
                f"ستون «{STOCK_HEADER}» در خروجی اکسل سایت پیدا نشد."
            )

        apply_inventory_row_colors(
            worksheet=worksheet,
            site_stock_column=stock_column,
            product_stock_column=product_stock_column,
            last_column=output_column,
        )

        workbook.save(output_path)

    finally:
        workbook.close()


def create_definition_date_report(
    site_file_path,
    product_stock_file_path,
    definition_file_path,
    output_path,
):
    output_path = Path(output_path)

    create_report(
        file1=site_file_path,
        file2=product_stock_file_path,
        output_path=output_path,
    )

    append_definition_dates(
        output_path=output_path,
        third_file_path=definition_file_path,
    )

    return str(output_path)
