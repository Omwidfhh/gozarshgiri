import shutil
from copy import copy
from math import isfinite
from pathlib import Path

import xlsxwriter
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from python_calamine import CalamineWorkbook


CODE_HEADER = "کد کالا"
STOCK_HEADER = "موجودی قابل فروش"

PRODUCT_STOCK_HEADER = "موجودی انبار محصول"
TEMP_STOCK_HEADER = "موجودی انبار موقت"


SITE_CHARGE_DROP_COLUMNS = (
    (
        "کالا نم انبار",
        "کالا نام انبار",
        "نام انبار کالا",
        "نام انبار",
    ),
    (
        "موجودی تعدادی",
        "موجودي تعدادي",
    ),
    (
        "موجودی رزرو کالا",
        "موجودي رزرو کالا",
    ),
    (
        "سایز",
        "سايز",
    ),
    (
        "مشخصات کالا",
        "مشخصات كالا",
    ),
    (
        "سایر مشخصات کالا",
        "ساير مشخصات کالا",
        "سایر مشخصات كالا",
    ),
    (
        "سایز کالا",
        "سايز کالا",
    ),
    (
        "category کالا",
        "categoryکالا",
        "دسته بندی کالا",
        "دسته‌بندی کالا",
    ),
)


LIGHT_ORANGE_FILL = PatternFill(
    fill_type="solid",
    fgColor="FCE4D6",
)

PEYDA_BODY_FONT = Font(
    name="Peyda",
    size=11,
)

PEYDA_HEADER_FONT = Font(
    name="Peyda",
    size=11,
    bold=True,
)


DIGIT_TRANSLATION = str.maketrans(
    "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
    "01234567890123456789",
)


DAILY_REPORT_COLUMNS = {
    "product_title": [
        "عنوان کالا",
    ],
    "technical_details": [
        "مشخصات فنی",
        "مشخصات فنی کالا",
        "مشخصه فنی",
        "مشخصه فنی کالا",
    ],
    "unit": [
        "واحد",
        "واحد سنجش",
        "عنوان واحد سنجش",
    ],
    "amount": [
        "مقدار",
    ],
    "issuer": [
        "صادر کننده",
        "صادرکننده",
    ],
}


DAILY_MINIMUM_WIDTHS = {
    1: 18,
    2: 18,
    3: 10,
    4: 10,
    5: 12,
    6: 14,
    7: 16,
}


def normalize_header(value):
    if value is None:
        return ""

    text = str(value)

    text = (
        text
        .replace("ي", "ی")
        .replace("ك", "ک")
        .replace("\u200c", " ")
        .replace("\u00a0", " ")
    )

    return " ".join(text.split())


def normalize_header_compact(value):
    return normalize_header(
        value
    ).replace(" ", "")


def normalize_code(value):
    if value is None:
        return None

    if isinstance(value, bool):
        return str(value)

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float):
        if not isfinite(value):
            return None

        if value.is_integer():
            return str(int(value))

        return format(value, ".15g")

    code = str(value).translate(
        DIGIT_TRANSLATION
    )

    code = (
        code
        .replace("\u200c", "")
        .replace("\u200e", "")
        .replace("\u200f", "")
        .replace("\u00a0", " ")
        .strip()
    )

    return code or None


def find_column(headers, column_name):
    normalized_name = normalize_header(
        column_name
    )

    for column_number, value in enumerate(
        headers,
        start=1,
    ):
        if normalize_header(value) == normalized_name:
            return column_number

    return None


def find_column_by_aliases(
    headers,
    aliases,
):
    normalized_aliases = {
        normalize_header_compact(alias)
        for alias in aliases
    }

    for column_index, header in enumerate(
        headers
    ):
        normalized_header_value = (
            normalize_header_compact(
                header
            )
        )

        if (
            normalized_header_value
            in normalized_aliases
        ):
            return column_index

    return None


def is_zero(value):
    if value is None:
        return False

    if isinstance(value, bool):
        return False

    if isinstance(value, (int, float)):
        return value == 0

    normalized_value = (
        str(value)
        .translate(DIGIT_TRANSLATION)
        .replace(",", "")
        .strip()
    )

    try:
        return float(normalized_value) == 0
    except ValueError:
        return False


def has_value(value):
    if value is None:
        return False

    if isinstance(value, str):
        return bool(value.strip())

    return True


def get_headers(sheet, file_label):
    try:
        return next(
            sheet.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True,
            )
        )

    except StopIteration as error:
        raise ValueError(
            f"{file_label} خالی است."
        ) from error


def read_stock_data(
    file_path,
    file_label,
):
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    sheet = workbook.active

    try:
        headers = get_headers(
            sheet,
            file_label,
        )

        code_column = find_column(
            headers,
            CODE_HEADER,
        )

        stock_column = find_column(
            headers,
            STOCK_HEADER,
        )

        if code_column is None:
            raise ValueError(
                f"ستون «کد کالا» در {file_label} پیدا نشد."
            )

        if stock_column is None:
            raise ValueError(
                f"ستون «موجودی قابل فروش» در {file_label} پیدا نشد."
            )

        stock_data = {}

        required_column = max(
            code_column,
            stock_column,
        )

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True,
        ):
            if len(row) < required_column:
                continue

            product_code = normalize_code(
                row[code_column - 1]
            )

            if product_code is None:
                continue

            stock_data[product_code] = (
                row[stock_column - 1]
            )

        return stock_data

    finally:
        workbook.close()


def apply_peyda_font(
    sheet,
    max_row,
    max_column,
):
    for row in sheet.iter_rows(
        min_row=1,
        max_row=max_row,
        min_col=1,
        max_col=max_column,
    ):
        for cell in row:
            if cell.row == 1:
                cell.font = PEYDA_HEADER_FONT
            else:
                cell.font = PEYDA_BODY_FONT


def delete_columns_by_aliases(
    sheet,
    alias_groups,
):
    if not alias_groups:
        return

    normalized_aliases = {
        normalize_header_compact(alias).casefold()
        for aliases in alias_groups
        for alias in aliases
    }

    columns_to_delete = []

    for cell in sheet[1]:
        normalized_header_value = (
            normalize_header_compact(
                cell.value
            ).casefold()
        )

        if normalized_header_value in normalized_aliases:
            columns_to_delete.append(
                cell.column
            )

    for column_number in sorted(
        columns_to_delete,
        reverse=True,
    ):
        sheet.delete_cols(
            column_number,
            1,
        )


def replace_nature_column_with_code(
    sheet,
):
    headers = get_headers(
        sheet,
        "اکسل سایت",
    )

    code_index = find_column_by_aliases(
        headers,
        (
            "کد کالا",
            "كد کالا",
            "کدکالا",
        ),
    )

    nature_index = find_column_by_aliases(
        headers,
        (
            "ماهیت کالا کالا",
            "ماهيت کالا کالا",
            "ماهیتکالاکالا",
            "ماهیت کالا",
            "ماهيت کالا",
            "ماهیتکالا",
        ),
    )

    if code_index is None:
        raise ValueError(
            "ستون «کد کالا» در اکسل سایت پیدا نشد."
        )

    if nature_index is None:
        raise ValueError(
            "ستون «ماهیت کالا کالا» در اکسل سایت پیدا نشد."
        )

    source_column = code_index + 1
    target_column = nature_index + 1

    if source_column == target_column:
        return

    merged_ranges_to_remove = [
        str(merged_range)
        for merged_range in list(
            sheet.merged_cells.ranges
        )
        if (
            merged_range.min_col
            <= source_column
            <= merged_range.max_col
        )
        or (
            merged_range.min_col
            <= target_column
            <= merged_range.max_col
        )
    ]

    for merged_range in merged_ranges_to_remove:
        sheet.unmerge_cells(
            merged_range
        )

    for row_number in range(
        1,
        sheet.max_row + 1,
    ):
        source_cell = sheet.cell(
            row=row_number,
            column=source_column,
        )

        target_cell = sheet.cell(
            row=row_number,
            column=target_column,
        )

        if isinstance(target_cell, MergedCell):
            continue

        target_cell.value = source_cell.value

        if source_cell.has_style:
            target_cell._style = copy(
                source_cell._style
            )

        if source_cell.number_format:
            target_cell.number_format = (
                source_cell.number_format
            )

        target_cell.alignment = copy(
            source_cell.alignment
        )

        target_cell.protection = copy(
            source_cell.protection
        )

    sheet.cell(
        row=1,
        column=target_column,
        value=CODE_HEADER,
    )

    sheet.delete_cols(
        source_column,
        1,
    )


def create_combined_report(
    base_file,
    stock_sources,
    output_path,
    drop_columns=None,
    replace_nature_with_code=False,
    highlight_zero_rows=True,
):
    base_file = Path(base_file)
    output_path = Path(output_path)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    shutil.copy2(
        base_file,
        output_path,
    )

    workbook = load_workbook(
        filename=output_path,
        read_only=False,
        data_only=False,
    )

    sheet = workbook.active

    try:
        base_headers = get_headers(
            sheet,
            "اکسل سایت",
        )

        base_code_column = find_column(
            base_headers,
            CODE_HEADER,
        )

        base_stock_column = find_column(
            base_headers,
            STOCK_HEADER,
        )

        if base_code_column is None:
            raise ValueError(
                "ستون «کد کالا» در اکسل سایت پیدا نشد."
            )

        if base_stock_column is None:
            raise ValueError(
                "ستون «موجودی قابل فروش» در اکسل سایت پیدا نشد."
            )

        loaded_sources = []

        for source in stock_sources:
            stock_data = read_stock_data(
                file_path=source["file"],
                file_label=source["label"],
            )

            loaded_sources.append(
                {
                    "title": source["title"],
                    "data": stock_data,
                }
            )

        last_header_column = max(
            (
                cell.column
                for cell in sheet[1]
                if cell.value not in (None, "")
            ),
            default=0,
        )

        output_columns = []

        for source in loaded_sources:
            existing_column = find_column(
                base_headers,
                source["title"],
            )

            if existing_column is not None:
                output_column = existing_column
            else:
                last_header_column += 1
                output_column = last_header_column

            sheet.cell(
                row=1,
                column=output_column,
                value=source["title"],
            )

            output_columns.append(
                {
                    "column": output_column,
                    "data": source["data"],
                }
            )

        final_column = max(
            last_header_column,
            *[
                item["column"]
                for item in output_columns
            ],
        )

        zero_rows = []

        for row_number in range(
            2,
            sheet.max_row + 1,
        ):
            product_code = normalize_code(
                sheet.cell(
                    row=row_number,
                    column=base_code_column,
                ).value
            )

            for output_item in output_columns:
                stock_value = output_item[
                    "data"
                ].get(
                    product_code,
                    "N/A",
                )

                sheet.cell(
                    row=row_number,
                    column=output_item["column"],
                    value=stock_value,
                )

            base_stock_value = sheet.cell(
                row=row_number,
                column=base_stock_column,
            ).value

            if (
                highlight_zero_rows
                and is_zero(base_stock_value)
            ):
                zero_rows.append(
                    row_number
                )

        for row_number in zero_rows:
            for column_number in range(
                1,
                final_column + 1,
            ):
                sheet.cell(
                    row=row_number,
                    column=column_number,
                ).fill = LIGHT_ORANGE_FILL

        if replace_nature_with_code:
            replace_nature_column_with_code(
                sheet
            )

        delete_columns_by_aliases(
            sheet=sheet,
            alias_groups=drop_columns,
        )

        final_column = sheet.max_column

        apply_peyda_font(
            sheet=sheet,
            max_row=sheet.max_row,
            max_column=final_column,
        )

        workbook.save(output_path)

    finally:
        workbook.close()

    return str(output_path)


def create_report(
    file1,
    file2,
    output_path="output/report.xlsx",
    highlight_zero_rows=True,
):
    return create_combined_report(
        base_file=file1,
        stock_sources=[
            {
                "file": file2,
                "label": "اکسل انبار محصول",
                "title": PRODUCT_STOCK_HEADER,
            }
        ],
        output_path=output_path,
        drop_columns=SITE_CHARGE_DROP_COLUMNS,
        replace_nature_with_code=True,
        highlight_zero_rows=highlight_zero_rows,
    )


def create_three_file_report(
    file1,
    file2,
    file3,
    output_path="output/report.xlsx",
):
    return create_combined_report(
        base_file=file1,
        stock_sources=[
            {
                "file": file2,
                "label": "اکسل انبار محصول",
                "title": PRODUCT_STOCK_HEADER,
            },
            {
                "file": file3,
                "label": "اکسل انبار موقت",
                "title": TEMP_STOCK_HEADER,
            },
        ],
        output_path=output_path,
        drop_columns=SITE_CHARGE_DROP_COLUMNS,
        replace_nature_with_code=True,
    )


def get_value_from_row(
    row,
    column_index,
):
    if column_index >= len(row):
        return None

    return row[column_index]


def find_daily_columns(
    headers,
    file_label,
):
    columns = {}

    readable_names = {
        "product_title":
            "عنوان کالا",

        "technical_details":
            "مشخصات فنی",

        "unit":
            "واحد",

        "amount":
            "مقدار",

        "issuer":
            "صادر کننده",
    }

    for key, aliases in (
        DAILY_REPORT_COLUMNS.items()
    ):
        column_index = (
            find_column_by_aliases(
                headers,
                aliases,
            )
        )

        if column_index is None:
            raise ValueError(
                f"ستون «{readable_names[key]}» "
                f"در {file_label} پیدا نشد."
            )

        columns[key] = column_index

    return columns


def calculate_visual_width(value):
    """
    محاسبه عرض تقریبی متن برای Excel.
    """
    if value is None:
        return 0

    text = str(value)

    width = 0

    for character in text:
        if (
            "\u0600" <= character <= "\u06ff"
        ):
            width += 1.25
        else:
            width += 1

    return width


def finalize_daily_report(
    output_path,
):
    """
    اجرای مرحله نهایی:

    - تنظیم واقعی فرمت Text ستون F
    - محاسبه عرض تمام ستون‌ها
    - اعمال فونت Peyda 11
    """
    workbook = load_workbook(
        filename=output_path,
        read_only=False,
        data_only=False,
    )

    sheet = workbook.active

    sheet.sheet_view.rightToLeft = True

    try:
        for column_number in range(1, 8):
            max_width = 0

            for row_number in range(
                1,
                sheet.max_row + 1,
            ):
                cell = sheet.cell(
                    row=row_number,
                    column=column_number,
                )

                if row_number == 1:
                    cell.font = PEYDA_HEADER_FONT
                else:
                    cell.font = PEYDA_BODY_FONT

                # فرمت واقعی Text برای کل ستون F
                if column_number == 6:
                    cell.number_format = "@"

                cell_width = calculate_visual_width(
                    cell.value
                )

                if cell_width > max_width:
                    max_width = cell_width

            minimum_width = (
                DAILY_MINIMUM_WIDTHS.get(
                    column_number,
                    10,
                )
            )

            calculated_width = max(
                minimum_width,
                max_width + 3,
            )

            calculated_width = min(
                calculated_width,
                60,
            )

            column_letter = get_column_letter(
                column_number
            )

            sheet.column_dimensions[
                column_letter
            ].width = calculated_width

            sheet.column_dimensions[
                column_letter
            ].bestFit = True

        workbook.save(output_path)

    finally:
        workbook.close()


def create_daily_load_report(
    files,
    output_path=(
        "output/daily-load-report.xlsx"
    ),
):
    """
    گزارش بار روز:

    A: عنوان کالا
    B: مشخصات فنی
    C: واحد
    D: ستون خالی
    E: مقدار
    F: حاصل‌ضرب D در E
    G: صادر کننده
    """
    if not files:
        raise ValueError(
            "حداقل یک سند انبار انتخاب کنید."
        )

    output_path = Path(output_path)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = xlsxwriter.Workbook(
        str(output_path),
        {
            "constant_memory": True,
            "nan_inf_to_errors": True,
            "strings_to_urls": False,
            "default_date_format":
                "yyyy-mm-dd hh:mm:ss",
        },
    )

    worksheet = workbook.add_worksheet(
        "Merged"
    )

    worksheet.right_to_left()
    worksheet.freeze_panes(1, 0)

    header_format = workbook.add_format(
        {
            "font_name": "Peyda",
            "font_size": 11,
            "bold": True,
            "bg_color": "#D9EAF7",
            "border": 1,
            "align": "center",
            "valign": "vcenter",
        }
    )

    body_format = workbook.add_format(
        {
            "font_name": "Peyda",
            "font_size": 11,
            "valign": "vcenter",
        }
    )

    text_formula_format = workbook.add_format(
        {
            "font_name": "Peyda",
            "font_size": 11,
            "num_format": "@",
            "valign": "vcenter",
        }
    )

    output_headers = [
        "عنوان کالا",
        "مشخصات فنی",
        "واحد",
        "",
        "مقدار",
        "",
        "صادر کننده",
    ]

    worksheet.write_row(
        0,
        0,
        output_headers,
        header_format,
    )

    output_row = 1

    try:
        for file_number, file_path in enumerate(
            files,
            start=1,
        ):
            file_label = (
                f"سند انبار {file_number}"
            )

            input_workbook = (
                CalamineWorkbook.from_path(
                    str(file_path)
                )
            )

            input_sheet = (
                input_workbook
                .get_sheet_by_index(0)
            )

            source_rows = input_sheet.to_python(
                skip_empty_area=False
            )

            if not source_rows:
                continue

            source_headers = source_rows[0]

            source_columns = find_daily_columns(
                source_headers,
                file_label,
            )

            for source_row in source_rows[1:]:
                product_title = get_value_from_row(
                    source_row,
                    source_columns[
                        "product_title"
                    ],
                )

                technical_details = (
                    get_value_from_row(
                        source_row,
                        source_columns[
                            "technical_details"
                        ],
                    )
                )

                unit = get_value_from_row(
                    source_row,
                    source_columns["unit"],
                )

                amount = get_value_from_row(
                    source_row,
                    source_columns["amount"],
                )

                issuer = get_value_from_row(
                    source_row,
                    source_columns["issuer"],
                )

                selected_values = [
                    product_title,
                    technical_details,
                    unit,
                    amount,
                    issuer,
                ]

                if not any(
                    has_value(value)
                    for value in selected_values
                ):
                    continue

                excel_row_number = (
                    output_row + 1
                )

                worksheet.write(
                    output_row,
                    0,
                    product_title,
                    body_format,
                )

                worksheet.write(
                    output_row,
                    1,
                    technical_details,
                    body_format,
                )

                worksheet.write(
                    output_row,
                    2,
                    unit,
                    body_format,
                )

                worksheet.write_blank(
                    output_row,
                    3,
                    None,
                    body_format,
                )

                worksheet.write(
                    output_row,
                    4,
                    amount,
                    body_format,
                )

                # فرمول تمام ردیف‌های ستون F
                worksheet.write_formula(
                    output_row,
                    5,
                    (
                        f"=D{excel_row_number}"
                        f"*E{excel_row_number}"
                    ),
                    text_formula_format,
                    0,
                )

                worksheet.write(
                    output_row,
                    6,
                    issuer,
                    body_format,
                )

                output_row += 1

        if output_row == 1:
            raise ValueError(
                "هیچ ردیف قابل پردازشی در فایل‌ها پیدا نشد."
            )

        worksheet.autofilter(
            0,
            0,
            output_row - 1,
            6,
        )

    finally:
        workbook.close()

    # مرحله دوم برای فرمت دقیق و تنظیم عرض‌ها
    finalize_daily_report(
        output_path
    )

    return str(output_path)
