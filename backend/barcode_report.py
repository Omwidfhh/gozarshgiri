from math import isfinite
from pathlib import Path

import xlsxwriter
from openpyxl import load_workbook


OUTPUT_HEADERS = (
    "عنوان کالا",
    "عنوان واحد سنجش",
    "مشخصه فنی کالا",
    "موجودی قابل فروش",
    "کد کالا",
    "کد کالا",
    "شناسه محصول",
)

BASE_COLUMN_ALIASES = {
    "product_title": (
        "عنوان کالا",
    ),
    "unit_title": (
        "عنوان واحد سنجش",
        "واحد سنجش",
        "واحد",
    ),
    "technical_code": (
        "مشخصه فنی کالا",
        "مشخصات فنی کالا",
        "مشخصه فنی",
        "مشخصات فنی",
    ),
    "sellable_stock": (
        "موجودی قابل فروش",
    ),
    "product_code": (
        "کد کالا",
    ),
}

BARCODE_ALIASES = (
    "بارکد",
    "بار کد",
    "کد بارکد",
)

THIRD_PRODUCT_CODE_ALIASES = (
    "کد محصول",
    "کد کالا",
)

PRODUCT_ID_ALIASES = (
    "شناسه",
    "شناسه محصول",
)

PERSIAN_DIGITS = str.maketrans(
    "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
    "01234567890123456789",
)


def has_value(value):
    if value is None:
        return False

    if isinstance(value, str):
        return bool(value.strip())

    return True


def normalize_text(value):
    if value is None:
        return ""

    return (
        str(value)
        .translate(PERSIAN_DIGITS)
        .replace("ي", "ی")
        .replace("ك", "ک")
        .replace("\u200c", " ")
        .replace("\u200e", "")
        .replace("\u200f", "")
        .replace("\xa0", " ")
        .strip()
    )


def normalize_header(value):
    return "".join(
        normalize_text(value).split()
    )


def normalize_code(value):
    if value is None:
        return ""

    if isinstance(value, bool):
        return str(value)

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float):
        if not isfinite(value):
            return ""

        if value.is_integer():
            return str(int(value))

        return format(value, ".15g")

    return "".join(
        normalize_text(value).split()
    )


def display_code(value):
    if value is None:
        return ""

    if isinstance(value, bool):
        return str(value)

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float):
        if not isfinite(value):
            return ""

        if value.is_integer():
            return str(int(value))

        return format(value, ".15g")

    return str(value).strip()


def get_headers(worksheet, file_label):
    try:
        return next(
            worksheet.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True,
            )
        )

    except StopIteration as error:
        raise ValueError(
            f"{file_label} خالی است."
        ) from error


def find_column(headers, aliases):
    normalized_aliases = {
        normalize_header(alias)
        for alias in aliases
    }

    for index, header in enumerate(headers):
        if normalize_header(header) in normalized_aliases:
            return index

    return None


def value_at(row, index):
    if index < len(row):
        return row[index]

    return None


def read_base_rows(file_path):
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook.active

    try:
        headers = get_headers(
            worksheet,
            "اکسل اول",
        )

        column_indexes = {}

        for key, aliases in BASE_COLUMN_ALIASES.items():
            column_index = find_column(
                headers,
                aliases,
            )

            if column_index is None:
                raise ValueError(
                    f"ستون «{aliases[0]}» در اکسل اول پیدا نشد."
                )

            column_indexes[key] = column_index

        rows = []

        for row in worksheet.iter_rows(
            min_row=2,
            values_only=True,
        ):
            selected_row = (
                value_at(row, column_indexes["product_title"]),
                value_at(row, column_indexes["unit_title"]),
                value_at(row, column_indexes["technical_code"]),
                value_at(row, column_indexes["sellable_stock"]),
                value_at(row, column_indexes["product_code"]),
            )

            if any(has_value(value) for value in selected_row):
                rows.append(selected_row)

        return rows

    finally:
        workbook.close()


def read_barcodes(file_path):
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook.active

    try:
        headers = get_headers(
            worksheet,
            "اکسل دوم",
        )

        barcode_index = find_column(
            headers,
            BARCODE_ALIASES,
        )

        if barcode_index is None:
            raise ValueError(
                "ستون «بارکد» در اکسل دوم پیدا نشد."
            )

        barcodes = []

        for row in worksheet.iter_rows(
            min_row=2,
            values_only=True,
        ):
            barcode = value_at(
                row,
                barcode_index,
            )

            if has_value(barcode):
                barcodes.append(barcode)

        return barcodes

    finally:
        workbook.close()


def read_product_ids(file_path):
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook.active

    try:
        headers = get_headers(
            worksheet,
            "اکسل سوم",
        )

        product_code_index = find_column(
            headers,
            THIRD_PRODUCT_CODE_ALIASES,
        )
        product_id_index = find_column(
            headers,
            PRODUCT_ID_ALIASES,
        )

        if product_code_index is None:
            raise ValueError(
                "ستون «کد محصول» در اکسل سوم پیدا نشد."
            )

        if product_id_index is None:
            raise ValueError(
                "ستون «شناسه» در اکسل سوم پیدا نشد."
            )

        product_ids = {}

        for row in worksheet.iter_rows(
            min_row=2,
            values_only=True,
        ):
            product_code = value_at(
                row,
                product_code_index,
            )
            product_key = normalize_code(
                product_code
            )

            if not product_key:
                continue

            product_id = value_at(
                row,
                product_id_index,
            )

            if (
                product_key not in product_ids
                or not has_value(product_ids[product_key])
            ):
                product_ids[product_key] = product_id

        return product_ids

    finally:
        workbook.close()


def write_regular_value(
    worksheet,
    row_number,
    column_number,
    value,
    cell_format,
):
    if value is None:
        worksheet.write_blank(
            row_number,
            column_number,
            None,
            cell_format,
        )
        return

    worksheet.write(
        row_number,
        column_number,
        value,
        cell_format,
    )


def create_barcode_comparison_report(
    base_file_path,
    barcode_file_path,
    product_file_path,
    output_path,
):
    base_file_path = Path(base_file_path)
    barcode_file_path = Path(barcode_file_path)
    product_file_path = Path(product_file_path)
    output_path = Path(output_path)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    base_rows = read_base_rows(
        base_file_path
    )

    barcodes = read_barcodes(
        barcode_file_path
    )

    product_ids = read_product_ids(
        product_file_path
    )

    technical_codes = {
        normalize_code(row[2])
        for row in base_rows
        if normalize_code(row[2])
    }

    barcode_codes = {
        normalize_code(barcode)
        for barcode in barcodes
        if normalize_code(barcode)
    }

    shared_codes = (
        technical_codes
        & barcode_codes
    )

    workbook = xlsxwriter.Workbook(
        str(output_path),
        {
            "constant_memory": True,
            "strings_to_numbers": False,
            "strings_to_urls": False,
        },
    )

    try:
        worksheet = workbook.add_worksheet(
            "مقایسه بارکد"
        )

        worksheet.right_to_left()
        worksheet.hide_gridlines(2)
        worksheet.freeze_panes(1, 0)

        header_format = workbook.add_format({
            "font_name": "Peyda",
            "font_size": 11,
            "bold": True,
            "font_color": "#FFFFFF",
            "bg_color": "#2F75B5",
            "align": "center",
            "valign": "vcenter",
            "border": 1,
            "border_color": "#D9E2F3",
        })

        text_format = workbook.add_format({
            "font_name": "Peyda",
            "font_size": 11,
            "valign": "vcenter",
        })

        center_format = workbook.add_format({
            "font_name": "Peyda",
            "font_size": 11,
            "align": "center",
            "valign": "vcenter",
        })

        code_format = workbook.add_format({
            "font_name": "Peyda",
            "font_size": 11,
            "align": "center",
            "valign": "vcenter",
            "num_format": "@",
        })

        highlighted_code_format = workbook.add_format({
            "font_name": "Peyda",
            "font_size": 11,
            "bold": True,
            "font_color": "#375623",
            "bg_color": "#E2F0D9",
            "align": "center",
            "valign": "vcenter",
            "num_format": "@",
        })

        for column_number, header in enumerate(
            OUTPUT_HEADERS
        ):
            worksheet.write(
                0,
                column_number,
                header,
                header_format,
            )

        worksheet.write_comment(
            0,
            4,
            "کد کالا از اکسل اول",
            {
                "author": "Royal Jeans",
            },
        )

        worksheet.write_comment(
            0,
            5,
            "بارکد واردشده از اکسل دوم",
            {
                "author": "Royal Jeans",
            },
        )

        worksheet.write_comment(
            0,
            6,
            "شناسه دریافت‌شده از اکسل سوم با تطبیق کد کالا و کد محصول",
            {
                "author": "Royal Jeans",
            },
        )

        output_row_count = max(
            len(base_rows),
            len(barcodes),
        )

        for row_index in range(output_row_count):
            excel_row = row_index + 1

            if row_index < len(base_rows):
                base_row = base_rows[row_index]

                write_regular_value(
                    worksheet,
                    excel_row,
                    0,
                    base_row[0],
                    text_format,
                )

                write_regular_value(
                    worksheet,
                    excel_row,
                    1,
                    base_row[1],
                    center_format,
                )

                technical_code = base_row[2]
                technical_key = normalize_code(
                    technical_code
                )

                technical_format = (
                    highlighted_code_format
                    if technical_key in shared_codes
                    else code_format
                )

                worksheet.write_string(
                    excel_row,
                    2,
                    display_code(technical_code),
                    technical_format,
                )

                write_regular_value(
                    worksheet,
                    excel_row,
                    3,
                    base_row[3],
                    center_format,
                )

                worksheet.write_string(
                    excel_row,
                    4,
                    display_code(base_row[4]),
                    code_format,
                )

                product_key = normalize_code(
                    base_row[4]
                )
                product_id = product_ids.get(
                    product_key
                )

                worksheet.write_string(
                    excel_row,
                    6,
                    display_code(product_id),
                    code_format,
                )

            if row_index < len(barcodes):
                barcode = barcodes[row_index]
                barcode_key = normalize_code(
                    barcode
                )

                barcode_format = (
                    highlighted_code_format
                    if barcode_key in shared_codes
                    else code_format
                )

                worksheet.write_string(
                    excel_row,
                    5,
                    display_code(barcode),
                    barcode_format,
                )

        worksheet.set_row(0, 25)
        worksheet.set_column(0, 0, 34)
        worksheet.set_column(1, 1, 20)
        worksheet.set_column(2, 2, 23)
        worksheet.set_column(3, 3, 19)
        worksheet.set_column(4, 6, 20)

        worksheet.autofilter(
            0,
            0,
            max(output_row_count, 1),
            len(OUTPUT_HEADERS) - 1,
        )

    finally:
        workbook.close()

    return str(output_path)
