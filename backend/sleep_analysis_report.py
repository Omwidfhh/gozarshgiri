import re
import shutil
from copy import copy
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from math import isfinite
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

from report import DIGIT_TRANSLATION
from report import PEYDA_BODY_FONT
from report import PEYDA_HEADER_FONT
from report import calculate_visual_width
from report import find_column_by_aliases
from report import normalize_code


SITE_CODE_ALIASES = (
    "کد کالا",
    "کدکالا",
    "كد کالا",
)

SITE_STOCK_ALIASES = (
    "موجودی قابل فروش",
    "موجودي قابل فروش",
)

TECHNICAL_CODE_ALIASES = (
    "مشخصه فنی کالا",
    "مشخصات فنی کالا",
    "مشخصه فنی",
    "مشخصات فنی",
)

PRODUCT_CODE_ALIASES = (
    "کد کالا",
    "کد محصول",
    "کدکالا",
    "کدمحصول",
)

PRODUCT_STOCK_ALIASES = (
    "موجودی قابل فروش",
    "موجودي قابل فروش",
)

SALES_CODE_ALIASES = (
    "کد محصول",
    "کدمحصول",
    "کد کالا",
)

SALES_QUANTITY_ALIASES = (
    "مقدار قلم سفارش خرده فروشی",
    "مقدار قلم سفارش خرده‌فروشی",
    "مقدار قلم سفارش خرده فروشي",
    "مقدار فروش",
)

BARCODE_ALIASES = (
    "بارکد",
    "بار کد",
    "کد بارکد",
)

DEFINITION_DATE_ALIASES = (
    "تاریخ تعریف",
    "تاريخ تعريف",
    "تاریخ ثبت",
    "تاريخ ثبت",
)

PRODUCT_STOCK_HEADER = "موجودی محصول"
SALES_QUANTITY_HEADER = "مقدار فروش"
TOTAL_STOCK_HEADER = "موجودی کل"
TURNOVER_RATE_HEADER = "نرخ گردش"
PRODUCT_AGE_HEADER = "سن محصول (روز)"
ANALYSIS_HEADER = "تحلیل خواب کالا"
DEFINITION_DATE_HEADER = "تاریخ تعریف"

FAST_STATUS = "پرفروش"
SLOW_STATUS = "کند فروش"
REVIEW_STATUS = "قابل بررسی"

SOFT_GREEN_FILL = PatternFill(
    fill_type="solid",
    fgColor="E2F0D9",
)

SOFT_RED_FILL = PatternFill(
    fill_type="solid",
    fgColor="F4CCCC",
)

SOFT_YELLOW_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFF2CC",
)

STATUS_FILLS = {
    FAST_STATUS: SOFT_GREEN_FILL,
    SLOW_STATUS: SOFT_RED_FILL,
    REVIEW_STATUS: SOFT_YELLOW_FILL,
}

DATE_PART_PATTERN = re.compile(r"\d+")

ANALYSIS_COMMENT = (
    "تحلیل ترکیبی بر اساس نرخ گردش، موجودی فعلی، مقدار فروش و سن محصول.\n"
    "پرفروش: فروش مناسب، نرخ گردش بالا یا پوشش موجودی کوتاه.\n"
    "کند فروش: محصول قدیمی با فروش صفر، نرخ گردش بسیار پایین یا پوشش موجودی طولانی.\n"
    "قابل بررسی: محصول تازه‌تعریف‌شده، تاریخ نامشخص یا وضعیت میانی."
)


def get_headers(worksheet, file_label):
    try:
        return next(
            worksheet.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True,
            )
        )
    except StopIteration as error:
        raise ValueError(
            f"{file_label} خالی است."
        ) from error


def find_required_column(
    headers,
    aliases,
    readable_name,
    file_label,
):
    column_index = find_column_by_aliases(
        headers,
        aliases,
    )

    if column_index is None:
        raise ValueError(
            f"ستون «{readable_name}» در {file_label} پیدا نشد."
        )

    return column_index + 1


def normalize_key(value):
    normalized = normalize_code(value)

    if normalized is None:
        return ""

    return "".join(
        str(normalized)
        .translate(DIGIT_TRANSLATION)
        .replace("ي", "ی")
        .replace("ك", "ک")
        .replace("\u200c", " ")
        .replace("\u200e", "")
        .replace("\u200f", "")
        .replace("\xa0", " ")
        .split()
    )


def to_decimal(value):
    if value is None or isinstance(value, bool):
        return None

    if isinstance(value, Decimal):
        return value

    if isinstance(value, int):
        return Decimal(value)

    if isinstance(value, float):
        if not isfinite(value):
            return None

        return Decimal(str(value))

    normalized = (
        str(value)
        .translate(DIGIT_TRANSLATION)
        .replace(",", "")
        .replace("٬", "")
        .replace("،", "")
        .replace(" ", "")
        .strip()
    )

    if normalized.casefold() in {
        "",
        "n/a",
        "na",
        "none",
        "null",
        "-",
    }:
        return None

    if (
        normalized.startswith("(")
        and normalized.endswith(")")
    ):
        normalized = (
            f"-{normalized[1:-1].strip()}"
        )

    try:
        return Decimal(normalized)
    except InvalidOperation:
        return None


def excel_number(value):
    number = to_decimal(value)

    if number is None:
        return 0

    if number == number.to_integral_value():
        return int(number)

    return float(number)


def is_gregorian_leap(year):
    return (
        year % 4 == 0
        and (
            year % 100 != 0
            or year % 400 == 0
        )
    )


def jalali_to_gregorian(year, month, day):
    jalali_year = year + 1595
    days = (
        -355668
        + (365 * jalali_year)
        + ((jalali_year // 33) * 8)
        + (((jalali_year % 33) + 3) // 4)
        + day
    )

    if month < 7:
        days += (month - 1) * 31
    else:
        days += ((month - 7) * 30) + 186

    gregorian_year = 400 * (days // 146097)
    days %= 146097

    if days > 36524:
        gregorian_year += 100 * (
            (days - 1) // 36524
        )
        days = (days - 1) % 36524

        if days >= 365:
            days += 1

    gregorian_year += 4 * (days // 1461)
    days %= 1461

    if days > 365:
        gregorian_year += (days - 1) // 365
        days = (days - 1) % 365

    gregorian_day = days + 1
    month_lengths = [
        0,
        31,
        29 if is_gregorian_leap(
            gregorian_year
        ) else 28,
        31,
        30,
        31,
        30,
        31,
        31,
        30,
        31,
        30,
        31,
    ]

    gregorian_month = 1

    while (
        gregorian_month <= 12
        and gregorian_day
        > month_lengths[gregorian_month]
    ):
        gregorian_day -= month_lengths[
            gregorian_month
        ]
        gregorian_month += 1

    return date(
        gregorian_year,
        gregorian_month,
        gregorian_day,
    )


def parse_definition_date(
    value,
    excel_epoch=None,
):
    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if (
        isinstance(value, (int, float))
        and not isinstance(value, bool)
        and isfinite(float(value))
        and 1 <= float(value) <= 100000
    ):
        try:
            converted = from_excel(
                value,
                epoch=excel_epoch,
            )
        except (
            TypeError,
            ValueError,
            OverflowError,
        ):
            converted = None

        if isinstance(converted, datetime):
            return converted.date()

        if isinstance(converted, date):
            return converted

    text = (
        str(value or "")
        .translate(DIGIT_TRANSLATION)
        .strip()
    )

    if text.casefold() in {
        "",
        "n/a",
        "na",
        "none",
        "null",
        "-",
    }:
        return None

    parts = [
        int(part)
        for part in DATE_PART_PATTERN.findall(text)
    ]

    if len(parts) == 1:
        compact = str(parts[0])

        if len(compact) == 8:
            parts = [
                int(compact[0:4]),
                int(compact[4:6]),
                int(compact[6:8]),
            ]

    if len(parts) < 3:
        return None

    first, second, third = parts[:3]

    if first >= 1200:
        year, month, day = first, second, third
    elif third >= 1200:
        year, month, day = third, second, first
    else:
        return None

    try:
        if year < 1700:
            return jalali_to_gregorian(
                year,
                month,
                day,
            )

        return date(year, month, day)
    except (ValueError, IndexError):
        return None


def read_product_stock(file_path):
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )
    worksheet = workbook.active

    try:
        headers = get_headers(
            worksheet,
            "اکسل محصول",
        )
        code_column = find_required_column(
            headers,
            PRODUCT_CODE_ALIASES,
            "کد کالا",
            "اکسل محصول",
        )
        stock_column = find_required_column(
            headers,
            PRODUCT_STOCK_ALIASES,
            "موجودی قابل فروش",
            "اکسل محصول",
        )
        required_column = max(
            code_column,
            stock_column,
        )
        product_stock = {}

        for row in worksheet.iter_rows(
            min_row=2,
            values_only=True,
        ):
            if len(row) < required_column:
                continue

            product_code = normalize_key(
                row[code_column - 1]
            )

            if not product_code:
                continue

            value = to_decimal(
                row[stock_column - 1]
            )

            product_stock[product_code] = (
                value
                if value is not None
                else Decimal("0")
            )

        return product_stock

    finally:
        workbook.close()


def read_sales_totals(file_path):
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )
    worksheet = workbook.active

    try:
        headers = get_headers(
            worksheet,
            "اکسل فروش",
        )
        code_column = find_required_column(
            headers,
            SALES_CODE_ALIASES,
            "کد محصول",
            "اکسل فروش",
        )
        quantity_column = find_required_column(
            headers,
            SALES_QUANTITY_ALIASES,
            "مقدار قلم سفارش خرده فروشی",
            "اکسل فروش",
        )
        required_column = max(
            code_column,
            quantity_column,
        )
        sales_totals = {}

        for row in worksheet.iter_rows(
            min_row=2,
            values_only=True,
        ):
            if len(row) < required_column:
                continue

            product_code = normalize_key(
                row[code_column - 1]
            )

            if not product_code:
                continue

            quantity = to_decimal(
                row[quantity_column - 1]
            )

            if quantity is None:
                quantity = Decimal("0")

            sales_totals[product_code] = (
                sales_totals.get(
                    product_code,
                    Decimal("0"),
                )
                + quantity
            )

        return sales_totals

    finally:
        workbook.close()


def read_definition_dates(file_path):
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )
    worksheet = workbook.active

    try:
        headers = get_headers(
            worksheet,
            "اکسل تاریخ تعریف",
        )
        barcode_column = find_required_column(
            headers,
            BARCODE_ALIASES,
            "بارکد",
            "اکسل تاریخ تعریف",
        )
        date_column = find_required_column(
            headers,
            DEFINITION_DATE_ALIASES,
            "تاریخ تعریف",
            "اکسل تاریخ تعریف",
        )
        required_column = max(
            barcode_column,
            date_column,
        )
        definition_dates = {}

        for row in worksheet.iter_rows(
            min_row=2,
            values_only=False,
        ):
            if len(row) < required_column:
                continue

            barcode = normalize_key(
                row[barcode_column - 1].value
            )

            if not barcode:
                continue

            date_cell = row[date_column - 1]
            parsed_date = parse_definition_date(
                date_cell.value,
                workbook.epoch,
            )
            candidate = {
                "value": date_cell.value,
                "number_format": (
                    date_cell.number_format
                ),
                "parsed_date": parsed_date,
            }
            existing = definition_dates.get(
                barcode
            )

            if existing is None:
                definition_dates[barcode] = candidate
                continue

            existing_date = existing[
                "parsed_date"
            ]

            if (
                parsed_date is not None
                and (
                    existing_date is None
                    or parsed_date < existing_date
                )
            ):
                definition_dates[barcode] = candidate

        return definition_dates

    finally:
        workbook.close()


def copy_cell_style(source_cell, target_cell):
    if source_cell.has_style:
        target_cell._style = copy(
            source_cell._style
        )

    target_cell.alignment = copy(
        source_cell.alignment
    )
    target_cell.border = copy(
        source_cell.border
    )
    target_cell.protection = copy(
        source_cell.protection
    )
    target_cell.number_format = (
        source_cell.number_format
    )


def get_last_header_column(worksheet):
    return max(
        (
            cell.column
            for cell in worksheet[1]
            if cell.value not in (None, "")
        ),
        default=worksheet.max_column,
    )


def ensure_output_column(
    worksheet,
    header,
    after_column,
):
    headers = get_headers(
        worksheet,
        "اکسل خروجی",
    )
    existing_index = find_column_by_aliases(
        headers,
        (header,),
    )

    if existing_index is not None:
        return existing_index + 1

    output_column = after_column + 1
    source_header = worksheet.cell(
        row=1,
        column=max(after_column, 1),
    )
    output_header = worksheet.cell(
        row=1,
        column=output_column,
    )

    if not isinstance(source_header, MergedCell):
        copy_cell_style(
            source_header,
            output_header,
        )

    output_header.value = header
    output_header.font = copy(
        PEYDA_HEADER_FONT
    )

    return output_column


def classify_product(
    site_stock,
    product_stock,
    sales_quantity,
    age_days,
):
    zero = Decimal("0")
    site_stock = max(site_stock, zero)
    product_stock = max(
        product_stock,
        zero,
    )
    sales_quantity = max(
        sales_quantity,
        zero,
    )
    total_stock = site_stock + product_stock
    total_supply = total_stock + sales_quantity

    turnover_rate = (
        sales_quantity / total_supply
        if total_supply > zero
        else zero
    )

    effective_age = max(
        age_days or 0,
        14,
    )
    daily_sales = (
        sales_quantity
        / Decimal(effective_age)
    )
    coverage_days = (
        total_stock / daily_sales
        if daily_sales > zero
        else None
    )

    if (
        sales_quantity > zero
        and total_stock == zero
    ):
        status = FAST_STATUS

    elif (
        age_days is not None
        and age_days < 21
    ):
        if (
            sales_quantity >= Decimal("3")
            and turnover_rate >= Decimal("0.55")
        ):
            status = FAST_STATUS
        else:
            status = REVIEW_STATUS

    elif (
        sales_quantity >= Decimal("3")
        and (
            turnover_rate >= Decimal("0.55")
            or (
                coverage_days is not None
                and coverage_days <= Decimal("45")
            )
        )
    ):
        status = FAST_STATUS

    elif (
        age_days is not None
        and age_days >= 21
        and total_stock > zero
        and sales_quantity == zero
    ):
        status = SLOW_STATUS

    elif (
        age_days is not None
        and age_days >= 45
        and total_stock > zero
        and (
            turnover_rate < Decimal("0.18")
            or (
                coverage_days is not None
                and coverage_days > Decimal("180")
            )
        )
    ):
        status = SLOW_STATUS

    elif (
        age_days is None
        and sales_quantity == zero
    ):
        status = REVIEW_STATUS

    elif (
        sales_quantity >= Decimal("2")
        and turnover_rate >= Decimal("0.45")
    ):
        status = FAST_STATUS

    else:
        status = REVIEW_STATUS

    return {
        "status": status,
        "total_stock": total_stock,
        "turnover_rate": turnover_rate,
        "coverage_days": coverage_days,
    }


def fill_analysis_row(
    worksheet,
    row_number,
    last_column,
    status,
):
    fill = STATUS_FILLS[status]

    for column_number in range(
        1,
        last_column + 1,
    ):
        cell = worksheet.cell(
            row=row_number,
            column=column_number,
        )

        if isinstance(cell, MergedCell):
            continue

        cell.fill = copy(fill)


def apply_output_font(
    worksheet,
    last_column,
    status_column,
):
    for row in worksheet.iter_rows(
        min_row=1,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=last_column,
    ):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue

            if cell.row == 1:
                cell.font = copy(
                    PEYDA_HEADER_FONT
                )
            elif cell.column == status_column:
                cell.font = Font(
                    name="Peyda",
                    size=11,
                    bold=True,
                )
            else:
                cell.font = copy(
                    PEYDA_BODY_FONT
                )


def set_output_widths(
    worksheet,
    output_columns,
):
    minimum_widths = {
        PRODUCT_STOCK_HEADER: 15,
        SALES_QUANTITY_HEADER: 14,
        TOTAL_STOCK_HEADER: 13,
        TURNOVER_RATE_HEADER: 12,
        PRODUCT_AGE_HEADER: 16,
        ANALYSIS_HEADER: 18,
        DEFINITION_DATE_HEADER: 16,
    }

    for header, column_number in output_columns.items():
        maximum_width = calculate_visual_width(
            header
        )

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):
            maximum_width = max(
                maximum_width,
                calculate_visual_width(
                    worksheet.cell(
                        row=row_number,
                        column=column_number,
                    ).value
                ),
            )

        column_letter = get_column_letter(
            column_number
        )
        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(
                maximum_width + 3,
                minimum_widths[header],
            ),
            28,
        )
        worksheet.column_dimensions[
            column_letter
        ].bestFit = True


def create_sleep_analysis_report(
    site_file_path,
    product_file_path,
    sales_file_path,
    definition_file_path,
    output_path,
):
    output_path = Path(output_path)
    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    product_stock = read_product_stock(
        product_file_path
    )
    sales_totals = read_sales_totals(
        sales_file_path
    )
    definition_dates = read_definition_dates(
        definition_file_path
    )

    shutil.copy2(
        site_file_path,
        output_path,
    )

    workbook = load_workbook(
        filename=output_path,
        read_only=False,
        data_only=False,
    )
    worksheet = workbook.active

    try:
        headers = get_headers(
            worksheet,
            "اکسل سایت",
        )
        site_code_column = find_required_column(
            headers,
            SITE_CODE_ALIASES,
            "کد کالا",
            "اکسل سایت",
        )
        site_stock_column = find_required_column(
            headers,
            SITE_STOCK_ALIASES,
            "موجودی قابل فروش",
            "اکسل سایت",
        )
        technical_column = find_required_column(
            headers,
            TECHNICAL_CODE_ALIASES,
            "مشخصه فنی کالا",
            "اکسل سایت",
        )

        last_column = get_last_header_column(
            worksheet
        )
        output_columns = {}

        for header in (
            PRODUCT_STOCK_HEADER,
            SALES_QUANTITY_HEADER,
            TOTAL_STOCK_HEADER,
            TURNOVER_RATE_HEADER,
            PRODUCT_AGE_HEADER,
            ANALYSIS_HEADER,
            DEFINITION_DATE_HEADER,
        ):
            output_column = ensure_output_column(
                worksheet=worksheet,
                header=header,
                after_column=last_column,
            )
            output_columns[header] = output_column
            last_column = max(
                last_column,
                output_column,
            )

        analysis_header_cell = worksheet.cell(
            row=1,
            column=output_columns[
                ANALYSIS_HEADER
            ],
        )
        analysis_header_cell.comment = Comment(
            ANALYSIS_COMMENT,
            "Royal Jeans",
        )

        today = date.today()
        counts = {
            FAST_STATUS: 0,
            SLOW_STATUS: 0,
            REVIEW_STATUS: 0,
        }
        definition_date_matches = 0

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):
            product_code = normalize_key(
                worksheet.cell(
                    row=row_number,
                    column=site_code_column,
                ).value
            )
            technical_code = normalize_key(
                worksheet.cell(
                    row=row_number,
                    column=technical_column,
                ).value
            )

            if not product_code:
                continue

            site_stock = to_decimal(
                worksheet.cell(
                    row=row_number,
                    column=site_stock_column,
                ).value
            ) or Decimal("0")
            matched_product_stock = (
                product_stock.get(
                    product_code,
                    Decimal("0"),
                )
            )
            matched_sales = sales_totals.get(
                product_code,
                Decimal("0"),
            )
            matched_definition = (
                definition_dates.get(
                    technical_code
                )
            )
            parsed_date = (
                matched_definition[
                    "parsed_date"
                ]
                if matched_definition is not None
                else None
            )
            age_days = (
                max((today - parsed_date).days, 0)
                if parsed_date is not None
                else None
            )
            analysis = classify_product(
                site_stock=site_stock,
                product_stock=(
                    matched_product_stock
                ),
                sales_quantity=matched_sales,
                age_days=age_days,
            )
            status = analysis["status"]
            counts[status] += 1

            row_values = {
                PRODUCT_STOCK_HEADER:
                    excel_number(
                        matched_product_stock
                    ),
                SALES_QUANTITY_HEADER:
                    excel_number(matched_sales),
                TOTAL_STOCK_HEADER:
                    excel_number(
                        analysis["total_stock"]
                    ),
                TURNOVER_RATE_HEADER:
                    float(
                        analysis["turnover_rate"]
                    ),
                PRODUCT_AGE_HEADER:
                    age_days,
                ANALYSIS_HEADER:
                    status,
            }

            for header, value in row_values.items():
                output_column = output_columns[
                    header
                ]
                source_cell = worksheet.cell(
                    row=row_number,
                    column=max(
                        output_column - 1,
                        1,
                    ),
                )
                output_cell = worksheet.cell(
                    row=row_number,
                    column=output_column,
                )

                if not isinstance(
                    source_cell,
                    MergedCell,
                ):
                    copy_cell_style(
                        source_cell,
                        output_cell,
                    )

                output_cell.value = value
                output_cell.font = copy(
                    PEYDA_BODY_FONT
                )

            worksheet.cell(
                row=row_number,
                column=output_columns[
                    TURNOVER_RATE_HEADER
                ],
            ).number_format = "0.0%"

            status_cell = worksheet.cell(
                row=row_number,
                column=output_columns[
                    ANALYSIS_HEADER
                ],
            )
            status_cell.font = Font(
                name="Peyda",
                size=11,
                bold=True,
            )

            definition_cell = worksheet.cell(
                row=row_number,
                column=output_columns[
                    DEFINITION_DATE_HEADER
                ],
            )
            adjacent_cell = worksheet.cell(
                row=row_number,
                column=max(
                    output_columns[
                        DEFINITION_DATE_HEADER
                    ] - 1,
                    1,
                ),
            )

            if not isinstance(
                adjacent_cell,
                MergedCell,
            ):
                copy_cell_style(
                    adjacent_cell,
                    definition_cell,
                )

            if matched_definition is not None:
                definition_date_matches += 1
                definition_cell.value = (
                    matched_definition["value"]
                )

                if matched_definition[
                    "number_format"
                ]:
                    definition_cell.number_format = (
                        matched_definition[
                            "number_format"
                        ]
                    )

            definition_cell.font = copy(
                PEYDA_BODY_FONT
            )

            fill_analysis_row(
                worksheet=worksheet,
                row_number=row_number,
                last_column=last_column,
                status=status,
            )

        apply_output_font(
            worksheet=worksheet,
            last_column=last_column,
            status_column=output_columns[
                ANALYSIS_HEADER
            ],
        )
        set_output_widths(
            worksheet=worksheet,
            output_columns=output_columns,
        )

        last_letter = get_column_letter(
            last_column
        )
        worksheet.auto_filter.ref = (
            f"A1:{last_letter}{worksheet.max_row}"
        )
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.rightToLeft = True

        workbook.save(output_path)

        return {
            "output_path": str(output_path),
            "row_count": max(
                worksheet.max_row - 1,
                0,
            ),
            "fast_selling_count": counts[
                FAST_STATUS
            ],
            "slow_selling_count": counts[
                SLOW_STATUS
            ],
            "review_count": counts[
                REVIEW_STATUS
            ],
            "definition_date_matched_count": (
                definition_date_matches
            ),
        }

    finally:
        workbook.close()
