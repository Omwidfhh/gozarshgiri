from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from math import isfinite
from pathlib import Path
import re

import xlsxwriter
from openpyxl import load_workbook
from xlsxwriter.utility import xl_range_abs, xl_rowcol_to_cell


FIRST_REQUIRED_COLUMNS = {
    "barcode": ("بارکد", "بار کد", "کد بارکد"),
    "stock": (
        "موجودی",
        "موجودی کالا",
        "موجودی قابل فروش",
        "موجودی انبار",
    ),
    "price": ("قیمت", "قیمت سایت", "قیمت فروش", "قیمت محصول"),
}

FIRST_DROP_COLUMNS = (
    ("توضیحات", "توضیحات محصول", "شرح"),
    ("تاریخ تعریف", "تاريخ تعريف"),
    ("تاریخ ویرایش", "تاريخ ويرايش", "تاریخ آخرین ویرایش"),
    ("شناسه راهکاران", "شناسه راهکار"),
)

SECOND_REQUIRED_COLUMNS = {
    "product_name": ("نام محصول", "عنوان محصول"),
    "technical": (
        "مشخصه فنی",
        "مشخصه فنی کالا",
        "مشخصات فنی",
        "مشخصات فنی کالا",
    ),
    "fee": ("فی", "في", "فی فروش", "فی صندوق", "قیمت فی"),
}

SECOND_DROP_COLUMNS = (
    ("کد محصول", "کد کالا"),
    ("قیمت مصرف کننده", "قيمت مصرف کننده"),
    ("عنوان اعلامیه", "عنوان اعلاميه"),
    ("از تاریخ", "از تاريخ"),
    ("تا تاریخ", "تا تاريخ"),
)

MATCHED_FEE_HEADER = "فی فراخوانی"
STATUS_HEADER = "وضعیت"
RAW_FEE_HEADER = "فی"
TECHNICAL_HEADER = "مشخصه فنی"

PERSIAN_DIGITS = str.maketrans(
    "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
    "01234567890123456789",
)
NUMERIC_WORD_PATTERN = re.compile(
    r"(?<!\w)عددی(?!\w)"
)
NUMBER_PATTERN = re.compile(
    r"^[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?$"
)


def has_value(value):
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    return True


def normalize_text(value):
    if value is None:
        return ""
    return (
        str(value)
        .translate(PERSIAN_DIGITS)
        .replace("ي", "ی")
        .replace("ك", "ک")
        .replace("ۀ", "ه")
        .replace("ة", "ه")
        .replace("\u200c", " ")
        .replace("\u200e", "")
        .replace("\u200f", "")
        .replace("\xa0", " ")
        .strip()
    )


def normalize_header(value):
    return re.sub(
        r"[\s_\-–—:/\\]+",
        "",
        normalize_text(value),
    ).casefold()


def normalize_code(value):
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not isfinite(value):
            return ""
        return str(int(value)) if value.is_integer() else format(value, ".15g")

    text = "".join(normalize_text(value).split())
    if not text:
        return ""

    numeric_text = (
        text.replace(",", "").replace("٬", "").replace("،", "")
    )
    if NUMBER_PATTERN.fullmatch(numeric_text):
        try:
            number = Decimal(numeric_text)
            if number == number.to_integral_value():
                return str(int(number))
        except InvalidOperation:
            pass
    return text.casefold()


def parse_number(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        return Decimal(str(value)) if isfinite(value) else None

    text = normalize_text(value)
    if not text:
        return None

    is_negative = text.startswith("(") and text.endswith(")")
    text = (
        text.replace("(", "")
        .replace(")", "")
        .replace(",", "")
        .replace("٬", "")
        .replace("،", "")
        .replace("٫", ".")
        .replace(" ", "")
    )
    match = re.search(
        r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?",
        text,
    )
    if not match:
        return None
    try:
        number = Decimal(match.group(0))
        return -number if is_negative and number > 0 else number
    except InvalidOperation:
        return None


def excel_number(value):
    if value is None:
        return None
    return int(value) if value == value.to_integral_value() else float(value)


def is_zero(value):
    number = parse_number(value)
    return number is not None and number == 0


def contains_numeric_word(value):
    if value is None:
        return False

    return bool(
        NUMERIC_WORD_PATTERN.search(
            normalize_text(value)
        )
    )


def value_at(row, index):
    return row[index] if index is not None and index < len(row) else None


def find_column(headers, aliases):
    normalized_aliases = {normalize_header(alias) for alias in aliases}
    for index, header in enumerate(headers):
        if normalize_header(header) in normalized_aliases:
            return index
    return None


def locate_header_row(worksheet, required_columns, file_label):
    maximum_row = min(worksheet.max_row or 1, 50)
    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=1,
            max_row=maximum_row,
            values_only=True,
        ),
        start=1,
    ):
        indexes = {
            key: find_column(row, aliases)
            for key, aliases in required_columns.items()
        }
        if all(index is not None for index in indexes.values()):
            nonempty_indexes = [
                index for index, header in enumerate(row) if has_value(header)
            ]
            last_index = max(nonempty_indexes, default=0)
            return row_number, tuple(row[: last_index + 1]), indexes

    required_text = "، ".join(
        aliases[0] for aliases in required_columns.values()
    )
    raise ValueError(
        f"ستون‌های موردنیاز در {file_label} پیدا نشدند. "
        f"ستون‌های لازم: {required_text}"
    )


def build_drop_indexes(headers, alias_groups):
    indexes = set()
    for aliases in alias_groups:
        index = find_column(headers, aliases)
        if index is not None:
            indexes.add(index)
    return indexes


def read_first_file(file_path):
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    worksheet = workbook.active
    try:
        header_row, headers, indexes = locate_header_row(
            worksheet,
            FIRST_REQUIRED_COLUMNS,
            "اکسل اول",
        )
        drop_indexes = build_drop_indexes(headers, FIRST_DROP_COLUMNS)
        kept_indexes = [
            index
            for index, header in enumerate(headers)
            if index not in drop_indexes and has_value(header)
        ]
        if indexes["price"] not in kept_indexes:
            raise ValueError("ستون «قیمت» نباید از خروجی اکسل اول حذف شود.")

        rows = []
        for row in worksheet.iter_rows(
            min_row=header_row + 1,
            values_only=True,
        ):
            if is_zero(value_at(row, indexes["stock"])):
                continue
            values = [value_at(row, index) for index in kept_indexes]
            barcode = value_at(row, indexes["barcode"])
            if any(has_value(value) for value in values) or has_value(barcode):
                rows.append({"values": values, "barcode": barcode})

        return {
            "headers": [headers[index] for index in kept_indexes],
            "price_output_index": kept_indexes.index(indexes["price"]),
            "barcode_output_index": kept_indexes.index(indexes["barcode"]),
            "rows": rows,
        }
    finally:
        workbook.close()


def read_second_file(file_path):
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    worksheet = workbook.active
    try:
        header_row, headers, indexes = locate_header_row(
            worksheet,
            SECOND_REQUIRED_COLUMNS,
            "اکسل دوم",
        )
        build_drop_indexes(headers, SECOND_DROP_COLUMNS)

        products = []
        for row in worksheet.iter_rows(
            min_row=header_row + 1,
            values_only=True,
        ):
            product_name = value_at(row, indexes["product_name"])
            if contains_numeric_word(product_name):
                continue

            technical = value_at(row, indexes["technical"])
            fee = value_at(row, indexes["fee"])
            if has_value(technical) or has_value(fee):
                products.append({
                    "technical": technical,
                    "fee": fee,
                })
        return products
    finally:
        workbook.close()


def values_are_equal(first_value, second_value):
    first_number = parse_number(first_value)
    second_number = parse_number(second_value)
    if first_number is not None and second_number is not None:
        return first_number == second_number

    first_text = normalize_text(first_value)
    second_text = normalize_text(second_value)
    return bool(first_text and second_text and first_text == second_text)


def calculate_visual_width(value):
    if value is None:
        return 0
    return sum(
        1.35 if "\u0600" <= character <= "\u06ff" else 1
        for character in str(value)
    )


def update_width(widths, column_index, value):
    widths[column_index] = max(
        widths[column_index],
        calculate_visual_width(value),
    )


def write_value(
    worksheet,
    row_index,
    column_index,
    value,
    text_format,
    center_format,
    date_format,
):
    if value is None:
        worksheet.write_blank(row_index, column_index, None, text_format)
    elif isinstance(value, (datetime, date)):
        worksheet.write_datetime(row_index, column_index, value, date_format)
    elif isinstance(value, bool):
        worksheet.write_boolean(row_index, column_index, value, center_format)
    elif isinstance(value, Decimal):
        worksheet.write_number(
            row_index,
            column_index,
            float(value),
            center_format,
        )
    else:
        worksheet.write(row_index, column_index, value, text_format)


def write_amount(worksheet, row_index, column_index, value, amount_format):
    number = parse_number(value)
    if number is not None:
        display_value = excel_number(number)
        worksheet.write_number(
            row_index,
            column_index,
            display_value,
            amount_format,
        )
        return display_value
    if value is None:
        worksheet.write_blank(row_index, column_index, None, amount_format)
        return None
    worksheet.write(row_index, column_index, value, amount_format)
    return value


def create_formats(workbook):
    header = workbook.add_format({
        "font_name": "Peyda",
        "font_size": 11,
        "bold": True,
        "font_color": "#FFFFFF",
        "bg_color": "#2F75B5",
        "align": "center",
        "valign": "vcenter",
        "border": 1,
        "border_color": "#D9E2F3",
    })
    base = {
        "font_name": "Peyda",
        "font_size": 11,
        "valign": "vcenter",
        "bottom": 1,
        "bottom_color": "#E7EAF0",
    }
    return {
        "header": header,
        "text": workbook.add_format({**base, "align": "right"}),
        "center": workbook.add_format({**base, "align": "center"}),
        "date": workbook.add_format({
            **base,
            "align": "center",
            "num_format": "yyyy-mm-dd",
        }),
        "amount": workbook.add_format({
            **base,
            "align": "center",
            "num_format": "#,##0.##",
        }),
        "status": workbook.add_format({
            **base,
            "bold": True,
            "align": "center",
        }),
        "ok": workbook.add_format({
            "font_name": "Peyda",
            "font_size": 11,
            "bold": True,
            "font_color": "#375623",
            "bg_color": "#E2F0D9",
            "align": "center",
            "valign": "vcenter",
        }),
        "no": workbook.add_format({
            "font_name": "Peyda",
            "font_size": 11,
            "bold": True,
            "font_color": "#9C0006",
            "bg_color": "#F4CCCC",
            "align": "center",
            "valign": "vcenter",
        }),
    }


def create_price_comparison_report(
    site_file_path,
    cashbox_file_path,
    output_path,
):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    first_file = read_first_file(Path(site_file_path))
    second_rows = read_second_file(Path(cashbox_file_path))
    base_headers = list(first_file["headers"])
    price_index = first_file["price_output_index"]
    barcode_base_index = first_file["barcode_output_index"]
    matched_fee_index = price_index + 1
    status_index = price_index + 2

    output_headers = list(base_headers)
    output_headers.insert(matched_fee_index, MATCHED_FEE_HEADER)
    output_headers.insert(status_index, STATUS_HEADER)
    raw_fee_index = len(output_headers)
    output_headers.append(RAW_FEE_HEADER)
    raw_technical_index = len(output_headers)
    output_headers.append(TECHNICAL_HEADER)

    barcode_index = (
        barcode_base_index + 2
        if barcode_base_index > price_index
        else barcode_base_index
    )

    second_lookup = {}
    for second_row in second_rows:
        technical_key = normalize_code(second_row["technical"])
        if technical_key and technical_key not in second_lookup:
            second_lookup[technical_key] = second_row["fee"]

    first_row_count = len(first_file["rows"])
    second_row_count = len(second_rows)
    output_row_count = max(first_row_count, second_row_count)

    workbook = xlsxwriter.Workbook(
        str(output_path),
        {
            "constant_memory": True,
            "strings_to_numbers": False,
            "strings_to_urls": False,
        },
    )
    try:
        worksheet = workbook.add_worksheet("بررسی قیمت محصولات")
        worksheet.right_to_left()
        worksheet.hide_gridlines(2)
        worksheet.freeze_panes(1, 0)
        worksheet.set_default_row(21)
        formats = create_formats(workbook)
        widths = [0.0] * len(output_headers)

        for column_index, header in enumerate(output_headers):
            worksheet.write(0, column_index, header, formats["header"])
            update_width(widths, column_index, header)

        raw_start_row = 1
        raw_end_row = max(second_row_count, 1)
        raw_fee_range = xl_range_abs(
            raw_start_row,
            raw_fee_index,
            raw_end_row,
            raw_fee_index,
        )
        raw_technical_range = xl_range_abs(
            raw_start_row,
            raw_technical_index,
            raw_end_row,
            raw_technical_index,
        )

        for output_row in range(1, output_row_count + 1):
            if output_row <= first_row_count:
                source_row = first_file["rows"][output_row - 1]
                base_values = list(source_row["values"])
                price_value = base_values[price_index]
                matched_fee = second_lookup.get(
                    normalize_code(source_row["barcode"])
                )

                output_values = list(base_values)
                output_values.insert(matched_fee_index, None)
                output_values.insert(status_index, None)

                for column_index, value in enumerate(output_values):
                    if column_index in (matched_fee_index, status_index):
                        continue
                    if column_index == price_index:
                        display_value = write_amount(
                            worksheet,
                            output_row,
                            column_index,
                            value,
                            formats["amount"],
                        )
                    elif column_index == barcode_index:
                        display_value = normalize_code(value)
                        worksheet.write_string(
                            output_row,
                            column_index,
                            display_value,
                            formats["text"],
                        )
                    else:
                        write_value(
                            worksheet,
                            output_row,
                            column_index,
                            value,
                            formats["text"],
                            formats["center"],
                            formats["date"],
                        )
                        display_value = value
                    update_width(widths, column_index, display_value)

                barcode_cell = xl_rowcol_to_cell(output_row, barcode_index)
                match_formula = (
                    f'=IFERROR(INDEX({raw_fee_range},'
                    f'MATCH({barcode_cell},{raw_technical_range},0)),"")'
                )
                worksheet.write_formula(
                    output_row,
                    matched_fee_index,
                    match_formula,
                    formats["amount"],
                    excel_number(parse_number(matched_fee))
                    if parse_number(matched_fee) is not None
                    else (matched_fee if matched_fee is not None else ""),
                )
                update_width(widths, matched_fee_index, matched_fee)

                status = (
                    "OK"
                    if values_are_equal(price_value, matched_fee)
                    else "NO"
                )
                price_cell = xl_rowcol_to_cell(output_row, price_index)
                matched_fee_cell = xl_rowcol_to_cell(
                    output_row,
                    matched_fee_index,
                )
                status_formula = (
                    f'=IF(OR({price_cell}="",{matched_fee_cell}=""),'
                    f'"NO",IF({price_cell}={matched_fee_cell},"OK","NO"))'
                )
                worksheet.write_formula(
                    output_row,
                    status_index,
                    status_formula,
                    formats["status"],
                    status,
                )
                update_width(widths, status_index, status)

            if output_row <= second_row_count:
                second_row = second_rows[output_row - 1]
                raw_fee = write_amount(
                    worksheet,
                    output_row,
                    raw_fee_index,
                    second_row["fee"],
                    formats["amount"],
                )
                technical_text = normalize_code(second_row["technical"])
                worksheet.write_string(
                    output_row,
                    raw_technical_index,
                    technical_text,
                    formats["text"],
                )
                update_width(widths, raw_fee_index, raw_fee)
                update_width(
                    widths,
                    raw_technical_index,
                    technical_text,
                )

        if first_row_count:
            for value, cell_format in (
                ("OK", formats["ok"]),
                ("NO", formats["no"]),
            ):
                worksheet.conditional_format(
                    1,
                    status_index,
                    first_row_count,
                    status_index,
                    {
                        "type": "text",
                        "criteria": "containing",
                        "value": value,
                        "format": cell_format,
                    },
                )

        worksheet.autofilter(
            0,
            0,
            output_row_count,
            len(output_headers) - 1,
        )
        worksheet.set_row(0, 25)
        for column_index, width in enumerate(widths):
            minimum = 15 if column_index in (
                price_index,
                matched_fee_index,
                status_index,
                raw_fee_index,
            ) else 13
            worksheet.set_column(
                column_index,
                column_index,
                min(max(width + 3, minimum), 45),
            )
    finally:
        workbook.close()

    return str(output_path)
